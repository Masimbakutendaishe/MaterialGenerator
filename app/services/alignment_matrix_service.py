"""Builds the Programme Alignment Matrix — a traceability table linking each learning
outcome to its assessment type and notional hours, matching real INSETA-style matrices."""
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.services.ai_service import generate_alignment_matrix_row
from app.services.document_service import _hex_to_rgb, _add_page_numbers, DEFAULT_PRIMARY

TYPE_LABELS = {
    "MC": "Multiple Choice", "SQ": "Short Question", "LQ": "Long Question",
    "ESS": "Essay", "WPA": "Workplace Application", "OTJ": "On The Job",
}


def build_alignment_matrix_docx(title: str, units: list, organization_name: str = None,
                                 seta: str = None, nqf_level: str = None, logo_bytes: bytes = None,
                                 brand_colors: dict = None, job_id: str = None) -> BytesIO:
    brand_colors = brand_colors or {}
    primary_hex = brand_colors.get("primary", DEFAULT_PRIMARY).lstrip("#") if brand_colors.get("primary") else DEFAULT_PRIMARY
    primary = _hex_to_rgb(primary_hex, DEFAULT_PRIMARY)

    doc = Document()

    from app.services.document_service import _build_branded_cover
    _build_branded_cover(doc, title, "Programme Alignment Matrix", organization_name, logo_bytes, primary, primary_hex, primary)

    # Programme info block
    info_table = doc.add_table(rows=0, cols=2)
    info_table.style = "Table Grid"
    info_rows = [("Programme Name", title)]
    if seta:
        info_rows.append(("SETA", seta))
    if nqf_level:
        info_rows.append(("NQF Level", nqf_level))
    info_rows.append(("Number of Units", str(len(units))))
    total_outcomes = sum(len(u.get("outcomes", [])) for u in units)
    info_rows.append(("Total Learning Outcomes", str(total_outcomes)))

    for label, value in info_rows:
        row = info_table.add_row().cells
        row[0].text = label
        row[0].paragraphs[0].runs[0].bold = True
        row[1].text = value

    doc.add_page_break()

    # Main matrix table
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Unit", "Learning Outcome", "Suggested Assessment Approach", "Notional Hours", "Ref"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(10)

    total_hours = 0
    for u_index, unit in enumerate(units, start=1):
        unit_name = unit.get("name", f"Unit {u_index}")
        outcomes = unit.get("outcomes", [])

        for o_index, outcome in enumerate(outcomes, start=1):
            row_data = generate_alignment_matrix_row(unit_name, outcome, course_title=title, job_id=job_id)
            assessment_type = row_data.get("assessment_type", "SQ")
            hours = row_data.get("notional_hours", 2)
            total_hours += hours

            row = table.add_row().cells
            row[0].text = unit_name if o_index == 1 else ""  # only show unit name once per group
            row[1].text = outcome
            row[2].text = f"{assessment_type} ({TYPE_LABELS.get(assessment_type, assessment_type)})"
            row[3].text = str(hours)
            row[4].text = f"{u_index}.{o_index}"

            for cell in row:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)

    doc.add_paragraph()
    total_p = doc.add_paragraph()
    total_run = total_p.add_run(f"Total Notional Hours: {total_hours}")
    total_run.bold = True

    # Legend
    doc.add_paragraph()
    legend_heading = doc.add_paragraph()
    legend_heading.add_run("Keys").bold = True
    for code, label in TYPE_LABELS.items():
        doc.add_paragraph(f"{code} = {label}", style="List Bullet")

    _add_page_numbers(doc)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer