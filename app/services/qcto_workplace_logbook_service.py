"""Builds the QCTO Workplace Logbook — Learner Details, a Personal Narrative reflection,
Witness Testimony, per-module 'About this Module' intro + a data-driven Scope of Work
Experience evidence table (from real work_experience_elements), the existing rich Daily
Work Log Entry blocks, a Contextualised Workplace Knowledge sign-off table, and a closing
Feedback/Judgement Report/Declaration — matching the real evidence-workbook structure used
in accredited QCTO workplace material."""
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.services.document_service import (
    _hex_to_rgb, _build_branded_cover, _add_branded_header_footer, _add_bottom_border,
    DEFAULT_PRIMARY, DEFAULT_SECONDARY,
)

# How many blank daily-entry pages to generate per module — enough for a real training period
ENTRIES_PER_MODULE = 10

def _add_wm_module_alignment_table(doc, primary, primary_hex, module):
    """Adds the per-module 'Module Alignment' table for a WM module, matching the
    reference pack structure. Unlike KM/PM, WM modules are judged Competent/Not Yet
    Competent rather than marked, so this shows notional hours instead of a mark total."""
    credits = module.get("credits", 0)
    try:
        credits = int(credits)
    except (ValueError, TypeError):
        credits = 0
    hours = credits * 10

    align_heading = doc.add_paragraph()
    ah_run = align_heading.add_run(f"{module.get('module_code', '')}  —  Module Alignment")
    ah_run.bold = True
    ah_run.font.size = Pt(13)
    ah_run.font.color.rgb = primary

    align_table = doc.add_table(rows=7, cols=2)
    align_table.style = "Table Grid"
    align_rows = [
        ("Work experience module", f"{module.get('module_code', '')}: {module.get('title', '')}"),
        ("Curriculum code", module.get("module_code", "")),
        ("NQF Level / Credits", f"{module.get('nqf_level', '')}  |  {credits}"),
        ("Notional workplace hours", f"{hours} hours"),
        ("Delivery weeks (per integrated register)", ""),
        ("Assessment methods", "Authenticated logbook; workplace evidence; coach or workplace SME observation and sign-off; witness statements"),
        ("Judgement", "Competent or Not Yet Competent. Work experience modules are not marked out of a total — every work activity must be evidenced and signed off."),
    ]
    for i, (label, value) in enumerate(align_rows):
        align_table.cell(i, 0).text = label
        align_table.cell(i, 0).paragraphs[0].runs[0].bold = True
        align_table.cell(i, 1).text = str(value)

    doc.add_paragraph()


def _add_wm_module_intro(doc, module, primary, primary_hex):
    """'About this Module' intro, data-driven from the module's real work_experience_elements
    and guidelines — matches the reference's compulsory-activity notice + WE requirement list."""
    heading = doc.add_paragraph()
    h_run = heading.add_run(f"About This Module: {module.get('title', '')}")
    h_run.bold = True
    h_run.font.size = Pt(16)
    h_run.font.color.rgb = primary
    _add_bottom_border(heading, primary_hex)

    notice_p = doc.add_paragraph()
    notice_run = notice_p.add_run(
        "The activities in this workbook are compulsory. You must complete them under supervision "
        "and they will be assessed by your facilitator. These activities must then be stored in "
        "your Portfolio of Evidence and submitted for assessment."
    )
    notice_run.bold = True

    points_heading = doc.add_paragraph()
    points_heading.add_run("Important Points to Remember:").italic = True
    points_heading.runs[0].bold = True
    for point in [
        "Complete the activities in your own words — do not copy from your learning guide.",
        "Please write legibly.",
        "Answer the questions in the space provided.",
        "Add additional pages if there is not enough space, and number them clearly.",
    ]:
        doc.add_paragraph(point, style="List Bullet")

    if module.get("guidelines"):
        purpose_p = doc.add_paragraph()
        purpose_p.add_run(module["guidelines"])
        purpose_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    elements = module.get("work_experience_elements") or []
    if elements:
        req_heading = doc.add_paragraph()
        req_heading.add_run("The learner will be required to:").bold = True
        for element in elements:
            doc.add_paragraph(element, style="List Bullet")

    doc.add_page_break()


def _add_personal_narrative(doc, primary, primary_hex):
    """Static reflection table — 'what went well / what would I do differently' — rendered
    once for the whole document, matching the reference's placement."""
    heading = doc.add_paragraph()
    h_run = heading.add_run("1. Personal Narrative")
    h_run.bold = True
    h_run.font.size = Pt(18)
    h_run.font.color.rgb = primary
    _add_bottom_border(heading, primary_hex)

    doc.add_paragraph(
        "Answer the following questions based on your experience during the completion of this "
        "module. Discuss what you did well and what you would like to do differently."
    )

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "What went well?"
    table.rows[0].cells[1].text = "What would I do differently?"
    for c in table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True

    prompts = [
        "I was able to identify and solve problems effectively throughout the various activities completed in this module.",
        "I was able to understand how different workplace activities have an impact on each other.",
        "I was able to use new technology effectively in my daily tasks.",
        "I was able to communicate effectively with my team members and supervisors.",
        "I was able to complete all my work in an organised and efficient manner.",
    ]
    for prompt in prompts:
        row = table.add_row().cells
        row[0].text = prompt
        row[0].paragraphs[0].runs[0].italic = True
        table.add_row()  # blank row for the learner to write in

    doc.add_paragraph()
    sig_table = doc.add_table(rows=2, cols=2)
    sig_table.style = "Table Grid"
    sig_table.cell(0, 0).text = "Learner Name:"
    sig_table.cell(0, 1).text = "Signature:"
    sig_table.cell(1, 0).text = "Date:"
    sig_table.cell(1, 1).text = ""
    doc.add_page_break()


def _add_witness_testimony(doc, primary, primary_hex):
    """Static supervisor/manager testimonial + multi-party acknowledgement sign-off,
    rendered once for the whole document."""
    heading = doc.add_paragraph()
    h_run = heading.add_run("2. Witness Testimony")
    h_run.bold = True
    h_run.font.size = Pt(18)
    h_run.font.color.rgb = primary
    _add_bottom_border(heading, primary_hex)

    doc.add_paragraph(
        "This section must be completed by the learner's supervisor / manager in the workplace, "
        "based on the learner's workplace performance relevant to the modules completed. "
        "Constructive comments and testimonial evidence may also be attached as a separate "
        "document and referenced below."
    )

    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Supervisor / Manager Testimonial"
    table.rows[0].cells[0].paragraphs[0].runs[0].bold = True
    for _ in range(3):
        table.add_row()

    doc.add_paragraph()
    for label in ["Supervisor Acknowledgement", "Assessor Acknowledgement", "Learner Acknowledgement", "Moderator Acknowledgement"]:
        ack_table = doc.add_table(rows=1, cols=3)
        ack_table.style = "Table Grid"
        ack_table.rows[0].cells[0].text = label
        ack_table.rows[0].cells[0].paragraphs[0].runs[0].bold = True
        ack_table.rows[0].cells[1].text = "Date:"
        ack_table.rows[0].cells[2].text = "Signature:"
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    doc.add_page_break()


def _add_scope_evidence_table(doc, module, primary, primary_hex):
    """Data-driven Scope of Work Experience table — one row per real work_experience_element,
    with Date/Signature sign-off columns, matching the WA-coded evidence table structure in
    the reference (without inventing sub-codes we don't have data for)."""
    elements = module.get("work_experience_elements") or []
    if not elements:
        return

    heading = doc.add_paragraph()
    h_run = heading.add_run("Scope of Work Experience")
    h_run.bold = True
    h_run.font.size = Pt(14)
    h_run.font.color.rgb = primary

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Work Experience Activity", "Date", "Signature"
    for c in hdr:
        c.paragraphs[0].runs[0].bold = True

    for i, element in enumerate(elements, start=1):
        row = table.add_row().cells
        row[0].text = f"{i}. {element}"

    doc.add_paragraph()


def _add_contextualised_knowledge_table(doc, primary, primary_hex):
    """Static contextualised workplace knowledge sign-off table, rendered once for the
    whole document, after all modules' evidence tables."""
    heading = doc.add_paragraph()
    h_run = heading.add_run("Contextualised Workplace Knowledge")
    h_run.bold = True
    h_run.font.size = Pt(16)
    h_run.font.color.rgb = primary
    _add_bottom_border(heading, primary_hex)

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Area", "Date", "Signature"
    for c in hdr:
        c.paragraphs[0].runs[0].bold = True

    for item in ["Workplace policies and procedures", "Workplace specific practices and customs",
                 "Workplace reporting structures", "Workplace documents"]:
        row = table.add_row().cells
        row[0].text = item

    doc.add_page_break()


def _add_judgement_and_declaration(doc, primary, primary_hex):
    """Closing feedback section, judgement report checklist, and learner declaration —
    the final assessment sign-off for the whole logbook."""
    heading = doc.add_paragraph()
    h_run = heading.add_run("Feedback Section")
    h_run.bold = True
    h_run.font.size = Pt(16)
    h_run.font.color.rgb = primary
    _add_bottom_border(heading, primary_hex)

    doc.add_paragraph("Comments from Learner:")
    for _ in range(3):
        doc.add_paragraph("_" * 100).paragraph_format.space_after = Pt(8)

    judgement_heading = doc.add_paragraph()
    jh_run = judgement_heading.add_run("Judgement Report")
    jh_run.bold = True
    jh_run.font.size = Pt(14)
    jh_run.font.color.rgb = primary

    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    outcomes = [
        ("Meets the requirements:", "Does not meet the requirements:"),
        ("Requires additional evidence:", "Requires another assessment:"),
        ("Can continue to the next assessment:", "Requires another assessment by another assessor:"),
    ]
    for i, (left, right) in enumerate(outcomes):
        table.cell(i, 0).text = f"☐ {left}"
        table.cell(i, 1).text = f"☐ {right}"

    doc.add_paragraph()
    doc.add_paragraph("Action required:")
    doc.add_paragraph("_" * 100)
    doc.add_paragraph("By when:")
    doc.add_paragraph("_" * 40)

    doc.add_page_break()

    decl_heading = doc.add_paragraph()
    dh_run = decl_heading.add_run("Declaration by Learner")
    dh_run.bold = True
    dh_run.font.size = Pt(16)
    dh_run.font.color.rgb = primary
    _add_bottom_border(decl_heading, primary_hex)

    doc.add_paragraph(
        "I, " + "_" * 40 + ", declare that I am satisfied that the feedback given to me by the "
        "Assessor was relevant, sufficient, and done in a constructive manner. I accept the "
        "assessment judgment and have no further questions relating to this assessment."
    )

    sig_table = doc.add_table(rows=2, cols=3)
    sig_table.style = "Table Grid"
    sig_table.cell(0, 0).text = "Learner Name & Signature"
    sig_table.cell(0, 1).text = "Assessor Name & Signature"
    sig_table.cell(0, 2).text = "Moderator Name & Signature"
    sig_table.cell(1, 0).text = "Date"
    sig_table.cell(1, 1).text = "Date"
    sig_table.cell(1, 2).text = "Date"


def _add_daily_log_entry(doc, primary_hex):
    """Adds one blank 'Daily Work Log Entry' block matching the reference format exactly."""
    date_p = doc.add_paragraph()
    date_p.add_run("Date: " + "_" * 50)
    date_p.paragraph_format.space_before = Pt(10)

    desc_heading = doc.add_paragraph()
    desc_run = desc_heading.add_run("Work Activity Description")
    desc_run.bold = True
    desc_sub = desc_heading.add_run(" (Write what tasks you performed today)")
    desc_sub.italic = True
    desc_sub.font.size = Pt(9)

    for _ in range(4):
        line_p = doc.add_paragraph()
        line_p.paragraph_format.space_after = Pt(6)
        line_p.add_run("_" * 100)

    for label in ["Activity Code(s) Performed:", "Location / Section of Work:", "Equipment / Software / Tool Used (if applicable):"]:
        field_p = doc.add_paragraph()
        field_p.add_run(f"{label} " + "_" * 50)
        field_p.paragraph_format.space_after = Pt(8)

    learnt_heading = doc.add_paragraph()
    learnt_heading.add_run("What I Learnt Today:").bold = True
    for _ in range(2):
        line_p = doc.add_paragraph()
        line_p.paragraph_format.space_after = Pt(6)
        line_p.add_run("_" * 100)

    query_heading = doc.add_paragraph()
    query_heading.add_run("Queries or Questions Related to Today's Activity:").bold = True
    for _ in range(2):
        line_p = doc.add_paragraph()
        line_p.paragraph_format.space_after = Pt(6)
        line_p.add_run("_" * 100)

    sig_heading = doc.add_paragraph()
    sig_heading.add_run("Signatures").bold = True
    sig_heading.paragraph_format.space_before = Pt(10)
    for label in ["Learner's Signature:", "Industrial Supervisor's Signature:", "Facilitator's Signature:"]:
        sig_p = doc.add_paragraph()
        sig_p.add_run(f"{label}  " + "_" * 40)
        sig_p.paragraph_format.space_after = Pt(6)

def _add_wm_document_control_and_roles(doc, primary, primary_hex, qualification_title):
    """Adds a Document Control table and 'A Note on Roles' section for the WM Logbook,
    matching the front-matter structure of a real QCTO work experience pack."""
    heading = doc.add_paragraph()
    h_run = heading.add_run("Document Control")
    h_run.bold = True
    h_run.font.size = Pt(14)
    h_run.font.color.rgb = primary
    _add_bottom_border(heading, primary_hex)

    control_table = doc.add_table(rows=7, cols=2)
    control_table.style = "Table Grid"
    control_fields = [
        ("Document title", f"Work Experience Logbook and Evidence Pack — {qualification_title}"),
        ("Document number", ""),
        ("Version", ""),
        ("Compiled by", ""),
        ("Approved by", ""),
        ("Date of issue", ""),
        ("Companion documents", ""),
    ]
    for i, (label, value) in enumerate(control_fields):
        control_table.cell(i, 0).text = label
        control_table.cell(i, 0).paragraphs[0].runs[0].bold = True
        control_table.cell(i, 1).text = value

    doc.add_paragraph()

    roles_heading = doc.add_paragraph()
    rh_run = roles_heading.add_run("A Note on Roles Used in This Pack")
    rh_run.bold = True
    rh_run.font.size = Pt(14)
    rh_run.font.color.rgb = primary
    _add_bottom_border(roles_heading, primary_hex)

    doc.add_paragraph(
        "The SME (Subject Matter Expert) is responsible for the programme and for preparing the "
        "learner. The coach or workplace SME is the experienced person who guides the learner on "
        "the job and signs off the logbook day by day. The final judgement of competence is made "
        "by a registered Assessor and verified by a registered Internal Moderator."
    )
    doc.add_page_break()


def _add_wm_overview_table(doc, primary, primary_hex, modules):
    """Adds the 'How the Work Experience Component Fits Together' overview, explaining
    that WM modules are judged Competent/Not Yet Competent rather than marked, plus a
    per-module hours breakdown table."""
    heading = doc.add_paragraph()
    h_run = heading.add_run("How the Work Experience Component Fits Together")
    h_run.bold = True
    h_run.font.size = Pt(14)
    h_run.font.color.rgb = primary
    _add_bottom_border(heading, primary_hex)

    doc.add_paragraph(
        "Work experience modules are not marked out of a total. The learner is judged Competent "
        "when every work activity has been performed, logged, evidenced and authenticated, and the "
        "target hours have been achieved. Completion of all work experience modules produces the "
        "Statement of Work Experience, which the QCTO requires before a learner may be entered for "
        "the External Integrated Summative Assessment."
    )

    table = doc.add_table(rows=len(modules) + 2, cols=4)
    table.style = "Table Grid"
    headers = ["Module", "Title", "Credits", "Hours"]
    for c, h in enumerate(headers):
        table.cell(0, c).text = h
        table.cell(0, c).paragraphs[0].runs[0].bold = True

    total_credits = 0
    total_hours = 0
    for i, module in enumerate(modules, start=1):
        credits = module.get("credits", 0)
        try:
            credits = int(credits)
        except (ValueError, TypeError):
            credits = 0
        hours = credits * 10
        total_credits += credits
        total_hours += hours

        table.cell(i, 0).text = module.get("module_code", "")
        table.cell(i, 1).text = module.get("title", "")
        table.cell(i, 2).text = str(credits)
        table.cell(i, 3).text = str(hours)

    total_row = len(modules) + 1
    table.cell(total_row, 0).text = "TOTAL"
    table.cell(total_row, 0).paragraphs[0].runs[0].bold = True
    table.cell(total_row, 2).text = str(total_credits)
    table.cell(total_row, 3).text = str(total_hours)

    doc.add_paragraph()
    doc.add_page_break()


def _add_wm_part_a(doc, primary, primary_hex, organization_name):
    """Adds Part A — Workplace Administration (A1-A4) for the WM Logbook: Learner and
    Workplace Details, Criteria for Workplace Approval (left as a template for the SDP
    to complete from their registered curriculum), Workplace Induction and Safety
    Record, and Workplace Agreement and Rules for completing the logbook."""
    part_heading = doc.add_paragraph()
    ph_run = part_heading.add_run("PART A — WORKPLACE ADMINISTRATION")
    ph_run.bold = True
    ph_run.font.size = Pt(16)
    ph_run.font.color.rgb = primary
    _add_bottom_border(part_heading, primary_hex)
    doc.add_paragraph()

    a1_heading = doc.add_paragraph()
    a1_run = a1_heading.add_run("A1  Learner and Workplace Details")
    a1_run.bold = True
    a1_run.font.size = Pt(13)
    a1_run.font.color.rgb = primary

    a1_fields = [
        "Surname and full names", "Learner / student number", "ID number", "Cell number",
        "Email address", "Employer", "Workplace where experience is gained",
        "Department / line", "Machines or equipment the learner works on", "Shift pattern",
        "Start date of workplace experience", "Expected end date",
    ]
    a1_table = doc.add_table(rows=len(a1_fields) + 1, cols=2)
    a1_table.style = "Table Grid"
    for i, field in enumerate(a1_fields):
        a1_table.cell(i, 0).text = field
        a1_table.cell(i, 0).paragraphs[0].runs[0].bold = True
    a1_table.cell(len(a1_fields), 0).text = "Employer"
    a1_table.cell(len(a1_fields), 0).paragraphs[0].runs[0].bold = True
    a1_table.cell(len(a1_fields), 1).text = organization_name or ""

    doc.add_paragraph()
    doc.add_paragraph("Coach and supervision").runs[0].bold = True
    coach_table = doc.add_table(rows=3, cols=3)
    coach_table.style = "Table Grid"
    coach_headers = ["Role", "Name", "Signature"]
    for c, h in enumerate(coach_headers):
        coach_table.cell(0, c).text = h
        coach_table.cell(0, c).paragraphs[0].runs[0].bold = True
    for i, role in enumerate(["Coach / Workplace SME", "SME (Subject Matter Expert)"], start=1):
        coach_table.cell(i, 0).text = role
    doc.add_page_break()

    a2_heading = doc.add_paragraph()
    a2_run = a2_heading.add_run("A2  Criteria for Workplace Approval (From the Registered Curriculum)")
    a2_run.bold = True
    a2_run.font.size = Pt(13)
    a2_run.font.color.rgb = primary

    doc.add_paragraph(
        "The curriculum sets these conditions for any workplace where this experience is gained. "
        "Complete each section from the qualification's actual registered curriculum document "
        "before this pack is issued."
    )
    for section_title in ["Physical requirements", "Human resource requirements", "Legal requirements"]:
        sub = doc.add_paragraph()
        sub.add_run(section_title).bold = True
        for _ in range(2):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run("_" * 80)

    doc.add_paragraph()
    doc.add_paragraph("Workplace approval confirmation").runs[0].bold = True
    approval_table = doc.add_table(rows=4, cols=2)
    approval_table.style = "Table Grid"
    approval_rows = [
        "The workplace meets the physical requirements above",
        "The learner works under the guidance of an experienced person",
        "The site complies with all applicable legal requirements",
        "The learner has access to the SOPs, work instructions and records this pack requires",
    ]
    for i, row in enumerate(approval_rows):
        approval_table.cell(i, 0).text = row
        approval_table.cell(i, 1).text = "Yes  /  No"
    doc.add_page_break()

    a3_heading = doc.add_paragraph()
    a3_run = a3_heading.add_run("A3  Workplace Induction and Safety Record")
    a3_run.bold = True
    a3_run.font.size = Pt(13)
    a3_run.font.color.rgb = primary

    doc.add_paragraph(
        "To be completed before the learner starts on the job. No work experience may be logged "
        "before this page is signed."
    )
    induction_items = [
        "Site health and safety induction",
        "Emergency procedures, assembly point and first aid",
        "PPE issued and correct use demonstrated",
        "Machine or equipment-specific safety briefing",
        "Site rules on reporting incidents, injuries and near misses",
        "Introduction to the coach and the team",
    ]
    induction_table = doc.add_table(rows=len(induction_items) + 1, cols=3)
    induction_table.style = "Table Grid"
    ind_headers = ["Induction item", "Date", "Conducted by"]
    for c, h in enumerate(ind_headers):
        induction_table.cell(0, c).text = h
        induction_table.cell(0, c).paragraphs[0].runs[0].bold = True
    for i, item in enumerate(induction_items, start=1):
        induction_table.cell(i, 0).text = item

    doc.add_paragraph()
    doc.add_paragraph(
        "Learner undertaking: I confirm that I received the induction recorded above, that I "
        "understand the safety rules that apply to me, and that I will work within them. I "
        "understand that I may not operate equipment I have not been briefed on, and that I must "
        "report any incident, injury or near miss immediately."
    )
    doc.add_paragraph("Learner: Signature: ________________  Date: ________________")
    doc.add_paragraph("Coach / Workplace SME: Signature: ________________  Date: ________________")
    doc.add_page_break()

    a4_heading = doc.add_paragraph()
    a4_run = a4_heading.add_run("A4  Workplace Agreement and Rules for Completing the Logbook")
    a4_run.bold = True
    a4_run.font.size = Pt(13)
    a4_run.font.color.rgb = primary

    doc.add_paragraph("What each party undertakes").runs[0].bold = True
    agreement_table = doc.add_table(rows=4, cols=2)
    agreement_table.style = "Table Grid"
    agreement_rows = [
        ("The learner", "Attends as scheduled; performs the work activities; makes logbook entries on the day, in ink, in own words; files the supporting evidence."),
        ("The coach / workplace SME", "Provides access to the equipment, documents and materials the module needs; guides and supervises the learner; verifies and signs each logbook entry."),
        ("The SME", "Prepares the learner; monitors progress against the module hours; collects and checks the pack."),
        ("The employer / SDP", "Maintains the workplace approval conditions; keeps the site compliant; retains the pack for moderation and verification."),
    ]
    for i, (party, undertaking) in enumerate(agreement_rows):
        agreement_table.cell(i, 0).text = party
        agreement_table.cell(i, 0).paragraphs[0].runs[0].bold = True
        agreement_table.cell(i, 1).text = undertaking

    doc.add_paragraph()
    doc.add_paragraph("Rules for completing the logbook").runs[0].bold = True
    rules = [
        "Write the entry on the day the work was done.",
        "Use black ink. Do not use correction fluid — draw a single line through an error, initial and date it.",
        "Write what you personally did. Copying the SOP into the logbook is not evidence.",
        "Record the location or line, the activity code, and the hours for every entry.",
        "Have your coach initial entries as you go, not at the end of the module.",
        "Attach the supporting evidence and mark each attachment with your name and the activity code.",
    ]
    for rule in rules:
        doc.add_paragraph(rule, style="List Bullet")
    doc.add_page_break()


def build_qcto_workplace_logbook_docx(title: str, syllabus_content: dict, organization_name: str = None,
                                       logo_bytes: bytes = None, brand_colors: dict = None,
                                       job_id: str = None) -> BytesIO:
    brand_colors = brand_colors or {}
    primary_hex = brand_colors.get("primary", DEFAULT_PRIMARY).lstrip("#") if brand_colors.get("primary") else DEFAULT_PRIMARY
    primary = _hex_to_rgb(primary_hex, DEFAULT_PRIMARY)
    qualification_title = syllabus_content.get("qualification_title", "") or title

    doc = Document()
    _build_branded_cover(doc, qualification_title, "WM Logbook", organization_name, logo_bytes, primary, primary_hex, primary)

    # Learner details table (front page, filled once)
    details_heading = doc.add_paragraph()
    details_heading.add_run("Learner Details (Fill on the first page only)").bold = True

    table = doc.add_table(rows=6, cols=2)
    table.style = "Table Grid"
    fields = ["Full Name", "Employee / Learner ID", "Department", "Supervisor Name",
              "Facilitator Name", "Logbook Start Date"]
    for i, field in enumerate(fields):
        table.cell(i, 0).text = field
        table.cell(i, 0).paragraphs[0].runs[0].bold = True

    doc.add_page_break()

    _add_wm_document_control_and_roles(doc, primary, primary_hex, qualification_title)

    modules = [m for m in syllabus_content.get("modules", []) if m.get("module_type") == "WM"]
    _add_wm_overview_table(doc, primary, primary_hex, modules)
    _add_wm_part_a(doc, primary, primary_hex, organization_name)

    _add_personal_narrative(doc, primary, primary_hex)
    _add_witness_testimony(doc, primary, primary_hex)

    # How to Use section
    how_heading = doc.add_paragraph()
    how_run = how_heading.add_run("How to Use This Logbook")
    how_run.bold = True
    how_run.font.size = Pt(16)
    how_run.font.color.rgb = primary
    _add_bottom_border(how_heading, primary_hex)

    how_text = doc.add_paragraph(
        "This logbook is designed to help you systematically record your daily work activities, "
        "learning experiences, and any challenges you encounter in the workplace — whether that "
        "involves machinery, software, tools, or other practical work, as applicable to your role. "
        "Keeping detailed and accurate records helps you track your progress, provides "
        "evidence of your hands-on experience for assessment purposes, and supports clear "
        "communication with your supervisors and facilitators."
    )
    how_text.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    instructions = [
        "Fill in your personal details in the Learner Details section on the first page.",
        "For each workday or shift, complete a new Daily Work Log Entry page.",
        "Record the date, tasks performed, relevant activity codes, and location/equipment, software, or tool used, as applicable.",
        "Reflect on what you learnt and note any questions for your supervisor or facilitator.",
        "Obtain signatures from yourself, your supervisor, and your facilitator for each entry.",
    ]
    for item in instructions:
        doc.add_paragraph(item, style="List Number")

    doc.add_page_break()

    wm_modules = [m for m in syllabus_content.get("modules", []) if m.get("module_type") == "WM"]

    for m_index, module in enumerate(wm_modules, start=1):
        _add_wm_module_alignment_table(doc, primary, primary_hex, module)
        _add_wm_module_intro(doc, module, primary, primary_hex)
        _add_scope_evidence_table(doc, module, primary, primary_hex)

        module_heading = doc.add_paragraph()
        mh_run = module_heading.add_run(f"Module {m_index} Logs: {module.get('title', '')}")
        mh_run.bold = True
        mh_run.font.size = Pt(16)
        mh_run.font.color.rgb = primary
        _add_bottom_border(module_heading, primary_hex)

        for entry_num in range(ENTRIES_PER_MODULE):
            _add_daily_log_entry(doc, primary_hex)
            if entry_num < ENTRIES_PER_MODULE - 1:
                doc.add_page_break()

        doc.add_page_break()

    _add_contextualised_knowledge_table(doc, primary, primary_hex)
    _add_judgement_and_declaration(doc, primary, primary_hex)

    _add_branded_header_footer(doc, logo_bytes=logo_bytes, qualification_name=qualification_title, organization_name=organization_name, primary_hex=primary_hex)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_qcto_workplace_logbook_docx_adapter(title, units, organization_name=None, seta=None,
                                               nqf_level=None, logo_bytes=None, brand_colors=None,
                                               job_id=None, **kwargs):
    """Adapter matching the standard DOCUMENT_BUILDERS call signature."""
    syllabus_content = {"modules": units} if isinstance(units, list) else (units or {"modules": []})
    return build_qcto_workplace_logbook_docx(
        title=title,
        syllabus_content=syllabus_content,
        organization_name=organization_name,
        logo_bytes=logo_bytes,
        brand_colors=brand_colors,
        job_id=job_id,
    )