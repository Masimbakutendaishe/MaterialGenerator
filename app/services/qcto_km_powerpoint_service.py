"""Builds the QCTO KM PowerPoint Presentations — one .pptx deck per real KM module,
bundled into a single ZIP file. Reuses generate_slide_content (the same AI function the
non-QCTO presentation builder uses) to write genuine teach/practice slide pairs per topic —
grounded in the topic's real title and assessment criteria, but actually written as
punchy, pedagogically-designed teaching content, not the raw criteria text reprinted as
bullets. Also reuses the app's established presentation house style from
presentation_service.py (split top bar, corner flag, footer band with organization name)
and adds page numbers to the footer."""
from io import BytesIO
import zipfile
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from app.services.ai_service import generate_slide_content, parallel_map
from app.services.document_service import DEFAULT_PRIMARY, DEFAULT_SECONDARY

DEFAULT_ACCENT = "F39C12"


def _pptx_rgb(hex_str, fallback):
    try:
        return RGBColor.from_string((hex_str or fallback).lstrip("#").upper())
    except (ValueError, TypeError):
        return RGBColor.from_string(fallback)


def _add_split_top_bar(slide, prs, color_left, color_right):
    """Two-tone top bar: left two-thirds in primary, right third in secondary — matches
    the house style used across every presentation this app generates."""
    bar_height = Emu(137160)
    split_point = int(prs.slide_width * 0.65)
    left_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), split_point, bar_height)
    left_bar.fill.solid()
    left_bar.fill.fore_color.rgb = color_left
    left_bar.line.fill.background()
    right_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, split_point, Emu(0), prs.slide_width - split_point, bar_height)
    right_bar.fill.solid()
    right_bar.fill.fore_color.rgb = color_right
    right_bar.line.fill.background()


def _add_corner_flag(slide, accent):
    flag = slide.shapes.add_shape(MSO_SHAPE.PENTAGON, Emu(0), Emu(137160), Inches(1.4), Inches(0.35))
    flag.fill.solid()
    flag.fill.fore_color.rgb = accent
    flag.line.fill.background()
    flag.rotation = 180


def _add_footer_band(slide, prs, organization_name, primary, page_number=None):
    """Solid corporate footer band — organization name on the left, page number on the
    right. Page numbers restart at 1 per module deck, since each is a standalone
    downloadable file a facilitator would present independently."""
    footer_height = Emu(320040)
    footer_top = prs.slide_height - footer_height
    footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), footer_top, prs.slide_width, footer_height)
    footer.fill.solid()
    footer.fill.fore_color.rgb = primary
    footer.line.fill.background()

    if organization_name:
        tf = footer.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_top = Emu(0)
        tf.margin_bottom = Emu(0)
        tf.paragraphs[0].text = organization_name
        run = tf.paragraphs[0].runs[0]
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    if page_number is not None:
        page_box = slide.shapes.add_textbox(prs.slide_width - Inches(1.2), footer_top, Inches(0.9), footer_height)
        p_tf = page_box.text_frame
        p_tf.margin_top = Emu(0)
        p_tf.margin_bottom = Emu(0)
        p_p = p_tf.paragraphs[0]
        p_p.alignment = PP_ALIGN.RIGHT
        p_run = p_p.add_run()
        p_run.text = str(page_number)
        p_run.font.size = Pt(10)
        p_run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _build_km_module_deck(module, qualification_title, organization_name, brand_colors, logo_bytes=None, job_id=None):
    primary = _pptx_rgb(brand_colors.get("primary"), DEFAULT_PRIMARY)
    secondary = _pptx_rgb(brand_colors.get("secondary"), DEFAULT_SECONDARY)
    accent = _pptx_rgb(brand_colors.get("accent"), DEFAULT_ACCENT)

    prs = Presentation()
    module_title = f"{module.get('module_code', '')}: {module.get('title', '')}"

    # --- Title slide ---
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    _add_split_top_bar(slide, prs, primary, secondary)
    _add_corner_flag(slide, accent)
    _add_footer_band(slide, prs, organization_name, primary, page_number=1)

    slide.shapes.title.text = module_title
    title_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    title_run.font.color.rgb = primary
    title_run.font.bold = True
    title_run.font.size = Pt(36)

    if qualification_title:
        subtitle_placeholder = slide.placeholders[1]
        subtitle_placeholder.text = qualification_title
        subtitle_placeholder.text_frame.paragraphs[0].runs[0].font.color.rgb = secondary
        subtitle_placeholder.text_frame.paragraphs[0].runs[0].font.italic = True
        subtitle_placeholder.text_frame.paragraphs[0].runs[0].font.size = Pt(18)

    if logo_bytes:
        try:
            slide.shapes.add_picture(BytesIO(logo_bytes), Inches(0.4), Inches(0.55), height=Inches(0.9))
        except Exception:
            pass  # unsupported image format for python-pptx (e.g. SVG) — skip rather than fail the deck

    # --- Content slides: real AI-generated teach/practice pairs per topic ---
    page_number = 2
    bullet_layout = prs.slide_layouts[1]

    topics = module.get("topics", [])
    topic_results = parallel_map(
        topics,
        lambda t: generate_slide_content(t.get("title", ""), t.get("assessment_criteria") or [], course_title=qualification_title, job_id=job_id),
        max_workers=3,
    )

    for topic, result in zip(topics, topic_results):
        unit_name = topic.get("title", "")
        if result is None:
            continue
        for slide_data in result.get("slides", []):
            slide_title = slide_data.get("title", unit_name)
            bullets = slide_data.get("bullets", [])
            speaker_notes = slide_data.get("speaker_notes", "")
            image_search_term = slide_data.get("image_search_term")

            photo_bytes = None
            if image_search_term:
                from app.services.image_service import fetch_stock_photo
                photo_bytes = fetch_stock_photo(image_search_term)

            slide = prs.slides.add_slide(bullet_layout)
            _add_split_top_bar(slide, prs, primary, secondary)
            _add_corner_flag(slide, accent)
            _add_footer_band(slide, prs, organization_name, primary, page_number=page_number)

            title_shape = slide.shapes.title
            title_shape.left = Inches(0.5)
            title_shape.top = Inches(1.0)
            title_shape.width = prs.slide_width - Inches(1.0)
            title_shape.height = Inches(1.0)
            title_shape.text = slide_title
            title_shape.text_frame.word_wrap = True
            title_shape.text_frame.paragraphs[0].runs[0].font.color.rgb = primary
            title_shape.text_frame.paragraphs[0].runs[0].font.bold = True

            # When a real photo was found, narrow the bullet area to make room for the
            # image alongside it rather than full-width — otherwise the bullets keep the
            # full slide width as before.
            content_width = prs.slide_width - Inches(1.0)
            body_width = int(content_width * 0.55) if photo_bytes else content_width

            body = slide.placeholders[1]
            body.left = Inches(0.5)
            body.top = Inches(2.2)
            body.width = body_width
            body.height = prs.slide_height - Inches(2.7)
            body.text_frame.clear()
            for i, bullet in enumerate(bullets):
                p = body.text_frame.paragraphs[0] if i == 0 else body.text_frame.add_paragraph()
                p.text = bullet
                p.font.size = Pt(20)

            if photo_bytes:
                image_left = Inches(0.5) + body_width + Inches(0.3)
                image_width = content_width - body_width - Inches(0.3)
                slide.shapes.add_picture(
                    BytesIO(photo_bytes), image_left, Inches(2.2),
                    width=image_width, height=prs.slide_height - Inches(2.7),
                )

            if speaker_notes:
                slide.notes_slide.notes_text_frame.text = speaker_notes

            page_number += 1

    buf = BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def build_qcto_km_powerpoint_zip(title: str, syllabus_content: dict, organization_name: str = None,
                                  logo_bytes: bytes = None, brand_colors: dict = None,
                                  job_id: str = None) -> BytesIO:
    brand_colors = brand_colors or {}
    qualification_title = syllabus_content.get("qualification_title", "") or title

    km_modules = [m for m in syllabus_content.get("modules", []) if m.get("module_type") == "KM"]

    zip_buf = BytesIO()
    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for module in km_modules:
            deck_buf = _build_km_module_deck(module, qualification_title, organization_name, brand_colors, logo_bytes, job_id=job_id)
            safe_title = "".join(c if c.isalnum() or c in " _-" else "" for c in module.get("title", "Module")).strip().replace(" ", "_")
            filename = f"{module.get('module_code', 'Module')}_{safe_title}.pptx"
            zf.writestr(filename, deck_buf.getvalue())

    zip_buf.seek(0)
    return zip_buf


def build_qcto_km_powerpoint_zip_adapter(title, units, organization_name=None, seta=None,
                                           nqf_level=None, logo_bytes=None, brand_colors=None,
                                           job_id=None, **kwargs):
    """Adapter matching the standard DOCUMENT_BUILDERS call signature."""
    syllabus_content = {"modules": units} if isinstance(units, list) else (units or {"modules": []})
    return build_qcto_km_powerpoint_zip(
        title=title,
        syllabus_content=syllabus_content,
        organization_name=organization_name,
        logo_bytes=logo_bytes,
        brand_colors=brand_colors,
        job_id=job_id,
    )
