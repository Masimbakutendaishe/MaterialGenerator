"""Portfolio of Evidence Guide — learner-facing document listing required evidence
per learning outcome, with space to record what was attached."""
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches
from app.services.ai_service import _call_model, _repair_json_string
from app.services.document_service import (
    _hex_to_rgb, _build_branded_cover, _add_signature_block, _add_branded_header_footer, DEFAULT_PRIMARY,
)
import json


def _generate_poe_item(unit_name: str, outcome: str, course_title: str = None, job_id: str = None) -> dict:
    """Generates a question, answer space, and evidence checklist for one learning outcome —
    a genuine learner-facing Portfolio of Evidence item, not just a table row."""
    prompt = f"""For this learning outcome from a South African SETA/QCTO-accredited programme,
write ONE Portfolio of Evidence item for the learner.

COURSE CONTEXT — this belongs to the course "{course_title or 'Unspecified Course'}". The question
and evidence examples MUST be genuinely relevant to that course's actual subject matter. If the unit
name or outcome is vague, interpret it strictly in the context of "{course_title}" — never substitute
in content from an unrelated field.

Unit: {unit_name}
Outcome: {outcome}

Write:
1. A question the learner must answer in their own words, demonstrating this outcome
2. How many ruled lines they'll need to answer it (2-5, based on expected answer length)
3. 1-2 pieces of REAL, concrete evidence they should collect and attach as proof (e.g. "Signed
   workplace observation checklist", "Copy of completed client needs analysis form")

Return ONLY valid JSON (no markdown, no commentary) in exactly this shape:
{{
  "question": "<the question text>",
  "blank_lines": 3,
  "evidence_examples": ["<evidence example 1>", "<evidence example 2>"]
}}"""

    raw = _call_model("slide_content", prompt, max_tokens=300, job_id=job_id)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        try:
            return json.loads(_repair_json_string(raw))
        except json.JSONDecodeError:
            return {"question": f"Describe how you demonstrated: {outcome}", "blank_lines": 3, "evidence_examples": ["Relevant workplace evidence as agreed with your facilitator"]}


def build_poe_guide_docx(title: str, units: list, organization_name: str = None,
                          seta: str = None, nqf_level: str = None, logo_bytes: bytes = None,
                          brand_colors: dict = None, job_id: str = None) -> BytesIO:
    brand_colors = brand_colors or {}
    primary_hex = brand_colors.get("primary", DEFAULT_PRIMARY).lstrip("#") if brand_colors.get("primary") else DEFAULT_PRIMARY
    primary = _hex_to_rgb(primary_hex, DEFAULT_PRIMARY)

    doc = Document()
    _build_branded_cover(doc, title, "Portfolio of Evidence", organization_name, logo_bytes, primary, primary_hex, primary)

    intro = doc.add_paragraph()
    intro.add_run(
        "This Portfolio of Evidence (PoE) requires you to answer questions demonstrating your "
        "achievement of each learning outcome, and to collect and attach real evidence as proof. "
        "Answer each question in your own words, then gather the evidence listed below it."
    )
    doc.add_paragraph()

    item_number = 1
    for u_index, unit in enumerate(units, start=1):
        unit_name = unit.get("name", f"Unit {u_index}")

        unit_heading = doc.add_paragraph()
        uh_run = unit_heading.add_run(unit_name)
        uh_run.bold = True
        uh_run.font.size = Pt(15)
        uh_run.font.color.rgb = primary

        for outcome in unit.get("outcomes", []):
            item = _generate_poe_item(unit_name, outcome, course_title=title, job_id=job_id)

            q_para = doc.add_paragraph()
            q_para.paragraph_format.space_before = Pt(10)
            q_para.add_run(f"{item_number}. {item.get('question', outcome)}").bold = True
            item_number += 1

            blank_lines = max(item.get("blank_lines", 3), min(round(item.get("marks", 0) * 0.8), 8))
            for _ in range(blank_lines):
                line_para = doc.add_paragraph()
                line_para.paragraph_format.space_after = Pt(12)
                line_para.add_run("_" * 100)

            evidence_heading = doc.add_paragraph()
            evidence_run = evidence_heading.add_run("Evidence to attach:")
            evidence_run.italic = True
            evidence_run.bold = True
            evidence_run.font.size = Pt(10)

            for evidence in item.get("evidence_examples", []):
                ev_p = doc.add_paragraph(f"☐ {evidence}", style="List Bullet")
                for run in ev_p.runs:
                    run.font.size = Pt(10)

            doc.add_paragraph()

        doc.add_page_break()

    # Summary evidence tracking table
    summary_heading = doc.add_paragraph()
    sh_run = summary_heading.add_run("Evidence Tracking Summary")
    sh_run.bold = True
    sh_run.font.size = Pt(16)
    sh_run.font.color.rgb = primary

    summary_intro = doc.add_paragraph()
    summary_intro.add_run(
        "Use this table to track which evidence you have collected and attached for each item."
    )

    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Item No.", "Learning Outcome", "Evidence Attached", "Page/Reference No."]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(10)

    summary_item_number = 1
    for unit in units:
        for outcome in unit.get("outcomes", []):
            row = table.add_row().cells
            row[0].text = str(summary_item_number)
            row[1].text = outcome
            row[2].text = ""
            row[3].text = ""
            for cell in row:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)
            summary_item_number += 1

    doc.add_page_break()
    _add_signature_block(doc)
    _add_branded_header_footer(doc, logo_bytes=logo_bytes, qualification_name=title, organization_name=organization_name, primary_hex=primary_hex)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer