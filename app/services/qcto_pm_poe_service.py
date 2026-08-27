"""Builds the QCTO PM Portfolio of Evidence (POE) — the fillable, learner-completed
companion to the PM Assessment Guide, mirroring the KM POE's structure but reframed for
PM's practical/scenario-based nature. Reuses the already-tested per-module scenario/PA-code/
IAC-code/Observation-Sheet section functions from qcto_pm_assessment_guide_service.py
(the learner needs the same practical activity content to actually complete, not a
duplicated re-implementation) rather than re-deriving that logic from scratch."""
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.services.document_service import (
    _hex_to_rgb, _build_branded_cover, _add_branded_header_footer, _add_bottom_border,
    DEFAULT_PRIMARY, DEFAULT_SECONDARY,
)
from app.services.qcto_pm_assessment_guide_service import (
    _add_pag_module_intro,
    _add_pag_scenario_and_activities,
)


def _section_heading(doc, text, primary, primary_hex, size=18):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = primary
    _add_bottom_border(p, primary_hex, size="8")
    return p


def _add_poe_cover_details(doc, qualification_title, primary, primary_hex):
    doc.add_paragraph(qualification_title).runs[0].bold = True
    table = doc.add_table(rows=8, cols=2)
    table.style = "Table Grid"
    fields = ["Learner Name", "Learner ID No.", "Company", "Contact No.",
              "Submission Date", "Attempt No.", "Assessor Name", "Assessor Reg No."]
    for i, field in enumerate(fields):
        table.cell(i, 0).text = field
        table.cell(i, 0).paragraphs[0].runs[0].bold = True
    doc.add_page_break()


def _add_poe_toc(doc, primary, primary_hex, module_count):
    _section_heading(doc, "Table of Contents", primary, primary_hex)
    entries = [
        "Foreword to Learner",
        "Assessment Process",
        "The Assessment Process Role-Players",
        "Competent and Not Yet Competent",
        "Requirements for Being Deemed Competent",
        "Note to the Learner",
        "Learner Personal Information",
        "Learner ID",
        "Pre-Assessment Preparation Sheet",
        "Assessment Plan",
    ]
    for i in range(1, module_count + 1):
        entries.append(f"Practical Applied Activities — Module {i}")
    entries += ["Workplace Evidence Diary", "Supervisor / Manager / Coach / Mentor Report", "Declaration of Authenticity"]
    for entry in entries:
        doc.add_paragraph(entry, style="List Bullet")
    doc.add_page_break()


def _add_poe_foreword(doc, primary, primary_hex):
    _section_heading(doc, "Foreword to Learner", primary, primary_hex)
    doc.add_paragraph(
        "The purpose of this guide is to provide you, the learner, with the process and "
        "requirements for successfully completing and submitting a Portfolio of Evidence "
        "for the practical component of this learning programme. Practical competence is "
        "demonstrated by doing, not by describing — the activities in this Portfolio ask "
        "you to actually perform each skill, observed directly against a realistic "
        "workplace scenario."
    )
    doc.add_paragraph(
        "Formative assessment refers to assessment that takes place during the process of "
        "learning and teaching, through supervised practice. Summative assessment is "
        "assessment for making a judgement about achievement, carried out through direct "
        "observation of a practical demonstration once you are ready. Results initially "
        "collected during formative practice can support summative assessment, avoiding "
        "unnecessary repetition."
    )
    doc.add_page_break()


def _add_poe_process(doc, primary, primary_hex):
    _section_heading(doc, "Assessment Process", primary, primary_hex)
    for i, step in enumerate([
        "Plan & prepare for assessment",
        "Conduct & record the assessment",
        "Provide feedback on the assessment",
        "Review and report on the assessment",
    ], start=1):
        doc.add_paragraph(f"Step {i}: {step}", style="List Number")
    doc.add_page_break()


def _add_poe_role_players(doc, primary, primary_hex, secondary):
    _section_heading(doc, "The Assessment Process Role-Players", primary, primary_hex)

    roles = [
        ("Learner", (
            "You will participate in demonstration and guided practice sessions using the "
            "PM Learner Guide and PM Facilitator Guide, before attempting each practical "
            "activity independently in the PM Assessment (Scenario-Based). You need to "
            "attend the practical sessions, practise under supervision, and complete all "
            "practical activities and portfolio requirements."
        )),
        ("Facilitator", (
            "The facilitator demonstrates each practical skill and guides your supervised "
            "practice using the PM Learner Guide and PM Facilitator Guide, is available for "
            "assessment-related questions, and guides you on the use of this POE."
        )),
        ("Assessor", (
            "The assessor must be qualified and registered with the relevant QCTO/SETA, "
            "proficient in the practical subject matter, and proficient in the assessment "
            "process — planning the assessment, observing your practical demonstration "
            "directly, providing feedback, recording results, and participating in "
            "moderation."
        )),
        ("Moderator", (
            "Internal moderators moderate assessment activities and support assessors, "
            "ensuring learners are assessed consistently and accurately across all "
            "assessors."
        )),
        ("Verifier", (
            "The moderation system is quality assured by the ETQAs, who have qualified "
            "verifiers in place to monitor moderation systems and support moderators."
        )),
        ("Training Provider", (
            "The training provider ensures qualified facilitators, assessors, and "
            "moderators are employed, and provides for an appeals process so learners can "
            "have an assessment reviewed if they believe the outcome was unfair."
        )),
    ]
    for name, desc in roles:
        heading = doc.add_paragraph()
        heading.add_run(name).bold = True
        heading.runs[0].font.color.rgb = secondary
        doc.add_paragraph(desc)
    doc.add_page_break()


def _add_poe_competence_and_evidence(doc, primary, primary_hex):
    _section_heading(doc, "Competent and Not Yet Competent", primary, primary_hex)
    doc.add_paragraph(
        "Practical modules are learned through demonstration, guided practice, and "
        "independent performance. Once you have demonstrated your competence through "
        "direct observation of a practical activity, you are awarded the credits related "
        "to that competence. If you are deemed not yet competent, you will either be given "
        "another chance to prove competence, be given further guided practice, or be "
        "encouraged to move into a different field of learning."
    )

    _section_heading(doc, "Requirements for Being Deemed Competent", primary, primary_hex, size=16)
    doc.add_paragraph(
        "Each practical module indicates the requirements or standards of competence, "
        "written as Internal Assessment Criteria. You need to meet ALL these requirements "
        "before being deemed competent. For practical skills, evidence is judged primarily "
        "through direct observation of your performance — watching you carry out the "
        "activity against a real workplace scenario — supplemented by any physical outputs "
        "you produce. Types of evidence include:"
    )
    for label, desc in [
        ("Direct", "Evidence collected directly by the assessor, such as observing you perform a task."),
        ("Indirect", "Evidence you have collected, signed off as authentic, and submitted for assessment."),
        ("Historic", "Evidence of your competence assessed by someone else, such as a prior certificate."),
    ]:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(desc)

    doc.add_paragraph()
    doc.add_paragraph("Evidence must meet the VARCS criteria:").runs[0].bold = True
    for letter, desc in [
        ("Valid", "The evidence submitted must genuinely be required by the assessment."),
        ("Authentic", "Evidence submitted must be your own work."),
        ("Relevant", "Evidence must relate directly to the assessment criteria."),
        ("Current", "Your evidence must demonstrate that your competence is current."),
        ("Sufficient", "The evidence must satisfy all the assessment criteria."),
    ]:
        p = doc.add_paragraph()
        p.add_run(f"{letter}: ").bold = True
        p.add_run(desc)
    doc.add_page_break()


def _add_poe_note_to_learner(doc, primary, primary_hex):
    _section_heading(doc, "Note to the Learner", primary, primary_hex)
    doc.add_paragraph(
        "Dear Learner, you have opted to undergo assessment and have been presented with "
        "this Portfolio of Evidence (POE). Please go through all sections very carefully "
        "before submission and make sure you have included all the information and "
        "evidence requested. Please take note of the following:"
    )
    for label, desc in [
        ("Pre-Assessment Preparation Sheet", "Read through this first — your assessor cannot assess your portfolio if you have not read and signed it."),
        ("Assessment Plan", "Use this to write down the dates on which you plan to meet specific targets."),
        ("Declaration of Authenticity", "Declare that the evidence you submit is your own work — your assessor cannot assess your portfolio without it."),
        ("Practical Applied Activities", "Complete each activity against the scenario given, following the instructions carefully."),
    ]:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(desc)
    doc.add_page_break()


def _add_poe_personal_info(doc, primary, primary_hex):
    _section_heading(doc, "Learner Personal Information", primary, primary_hex)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for field in ["Last Name", "First Name(s)", "Title", "Date of Birth", "ID Number",
                  "Nationality", "Gender", "Home Language", "Disability", "Home Address",
                  "Postal Address", "Telephone", "Cell Phone", "E-mail", "Date of Submission"]:
        row = table.add_row().cells
        row[0].text = field
    table.rows[0].cells[0].text = "Field"
    table.rows[0].cells[1].text = "Details"
    for c in table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True
    doc.add_page_break()

    _section_heading(doc, "Learner ID", primary, primary_hex)
    doc.add_paragraph("Insert a certified copy of your Identity Document here.")
    doc.add_page_break()


def _add_poe_preparation_sheet(doc, primary, primary_hex):
    _section_heading(doc, "Pre-Assessment Preparation Sheet", primary, primary_hex)
    doc.add_paragraph(
        "This document serves to orientate and prepare you for the practical assessment(s) "
        "you are about to undergo. It informs you of the steps involved and helps you "
        "prepare, giving you the best opportunity for success. This document MUST be "
        "completed in the presence of the Assessor/Facilitator."
    )

    info_table = doc.add_table(rows=1, cols=2)
    info_table.style = "Table Grid"
    for field in ["Venue / Pre-Assessment Meeting", "Date", "Learner Full Name", "Learner ID",
                  "Assessor Name", "Assessor Contact", "Assessor Number",
                  "Moderator Name", "Moderator Contact", "Moderator Number"]:
        row = info_table.add_row().cells
        row[0].text = field
    info_table.rows[0].cells[0].text = "Field"
    info_table.rows[0].cells[1].text = "Details"
    for c in info_table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True

    doc.add_paragraph()
    doc.add_paragraph("Please read the discussion points below and indicate that you have read and understand the information provided:").runs[0].bold = True

    discuss_table = doc.add_table(rows=1, cols=3)
    discuss_table.style = "Table Grid"
    discuss_table.rows[0].cells[0].text = "Discussion Point"
    discuss_table.rows[0].cells[1].text = "Yes / No"
    discuss_table.rows[0].cells[2].text = "Comments"
    for c in discuss_table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True

    for point in [
        "Were you welcomed and made to feel at ease?",
        "Was the purpose and objectives of the meeting explained?",
        "Was the practical assessment process and principles of good assessment explained?",
        "The purpose of the assessment is to determine and recognise your practical competence against the modules in this qualification.",
        "Do you understand the roles and responsibilities of the learner, assessor, and moderator?",
        "Were you informed of your rights, the appeal process, and re-assessment policies?",
        "Do you understand how and when the assessor will provide feedback?",
        "Do you understand the re-assessment opportunity available to you if evidence requirements are not met?",
        "Do you understand the recordkeeping and reporting of your results?",
    ]:
        row = discuss_table.add_row().cells
        row[0].text = point

    doc.add_paragraph()
    doc.add_paragraph("Declaration of Understanding:").runs[0].bold = True
    for statement in [
        "I understand the importance of this meeting.",
        "I declare that the above points were explained by the Assessor/Facilitator.",
        "I declare that I have received copies of the curriculum, assessment plan, and relevant policies pertaining to my assessment.",
        "I have read the above and understood the contents thereof.",
        "I was given the opportunity to clarify any issues relating to the assessment process and my assessment plan.",
        "I have requested this assessment in accordance with my own free will and without duress.",
    ]:
        doc.add_paragraph(statement, style="List Bullet")

    doc.add_paragraph()
    sig_table = doc.add_table(rows=4, cols=2)
    sig_table.style = "Table Grid"
    for i, label in enumerate(["Learner Signature", "Facilitator Signature", "Assessor Signature", "Moderator Signature"]):
        sig_table.cell(i, 0).text = label
        sig_table.cell(i, 1).text = "Date"
        sig_table.cell(i, 0).paragraphs[0].runs[0].bold = True
        sig_table.cell(i, 1).paragraphs[0].runs[0].bold = True
    doc.add_page_break()


def _add_poe_assessment_plan(doc, primary, primary_hex):
    _section_heading(doc, "Assessment Plan", primary, primary_hex)
    doc.add_paragraph("Use this plan to write down the dates on which you plan to meet specific targets. This document MUST be completed in the presence of the Assessor/Facilitator.")

    info_table = doc.add_table(rows=1, cols=2)
    info_table.style = "Table Grid"
    for field in ["Learner Full Name", "Learner ID", "Facilitator Name", "Assessor Name", "Assessor Number"]:
        row = info_table.add_row().cells
        row[0].text = field
    info_table.rows[0].cells[0].text = "Field"
    info_table.rows[0].cells[1].text = "Details"
    for c in info_table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True

    doc.add_paragraph()
    plan_table = doc.add_table(rows=1, cols=4)
    plan_table.style = "Table Grid"
    plan_table.rows[0].cells[0].text = "Action"
    plan_table.rows[0].cells[1].text = "Planned Date"
    plan_table.rows[0].cells[2].text = "Actual Date"
    plan_table.rows[0].cells[3].text = "Comments"
    for c in plan_table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True

    for action in [
        "Read and sign the Pre-Assessment Preparation Sheet",
        "Complete guided practice for each practical skill",
        "Complete the Practical Applied Activities in the PM Assessment (Scenario-Based)",
        "Complete the Learner's Review of the Assessment Process",
        "Submit the Portfolio of Evidence",
    ]:
        row = plan_table.add_row().cells
        row[0].text = action

    doc.add_paragraph()
    doc.add_paragraph(
        "I, the learner, hereby agree to the above plan and commit to preparing for the "
        "assessment and submitting the specified documents on the dates specified."
    )
    sig_table = doc.add_table(rows=3, cols=2)
    sig_table.style = "Table Grid"
    for i, label in enumerate(["Learner Signature", "Facilitator Signature", "Assessor Signature"]):
        sig_table.cell(i, 0).text = label
        sig_table.cell(i, 1).text = "Date"
        sig_table.cell(i, 0).paragraphs[0].runs[0].bold = True
        sig_table.cell(i, 1).paragraphs[0].runs[0].bold = True
    doc.add_page_break()


def _add_poe_workplace_diary(doc, primary, primary_hex):
    _section_heading(doc, "Workplace Evidence Diary", primary, primary_hex)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Date"
    table.rows[0].cells[1].text = "Diary Entry of Workplace Evidence"
    table.rows[0].cells[2].text = "Evidence Locator — Note Applicable Module"
    for c in table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True
    for _ in range(10):
        table.add_row()
    doc.add_page_break()


def _add_poe_supervisor_report(doc, primary, primary_hex):
    _section_heading(doc, "Supervisor / Manager / Coach / Mentor Report", primary, primary_hex)
    doc.add_paragraph(
        "Confidential. As part of the assessment for these practical modules, we are "
        "seeking evidence to support a judgement about the candidate's competence, "
        "including reports from the manager and other people who work closely with the "
        "candidate. Please submit this form to the Training Provider."
    )
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for field in ["Name of Candidate", "Learner Number", "Module(s) of Competency",
                  "Name of Supervisor / Manager / Coach / Mentor", "Workplace", "Address", "Phone"]:
        row = table.add_row().cells
        row[0].text = field
    table.rows[0].cells[0].text = "Field"
    table.rows[0].cells[1].text = "Details"
    for c in table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True
    doc.add_paragraph()
    doc.add_paragraph("Comments on candidate's practical competence:").runs[0].bold = True
    for _ in range(3):
        doc.add_paragraph("_" * 100)
    doc.add_page_break()


def _add_poe_declaration(doc, qualification_title, primary, primary_hex):
    _section_heading(doc, "Declaration of Authenticity", primary, primary_hex)
    doc.add_paragraph(
        "I, " + "_" * 40 + " (full name and surname), ID number " + "_" * 20 + ", declare "
        f"that the evidence presented in this Portfolio of Evidence represents workplace and "
        f"training evidence against the Programme: {qualification_title}, and that it is my "
        "own work, with the exception of:"
    )
    doc.add_paragraph("_" * 100)
    doc.add_paragraph()
    doc.add_paragraph("Please list any references to resources used (books, websites, etc.):")
    doc.add_paragraph("_" * 100)
    doc.add_paragraph()
    doc.add_paragraph(
        "In signing this, I declare that all the evidence presented in this Portfolio of "
        "Evidence is true, valid, and my own work."
    )

    sig_table = doc.add_table(rows=5, cols=2)
    sig_table.style = "Table Grid"
    for i, label in enumerate(["Learner Signature", "Witness Name", "Witness Signature", "Assessor Signature", "Moderator Signature"]):
        sig_table.cell(i, 0).text = label
        sig_table.cell(i, 1).text = "Date"
        sig_table.cell(i, 0).paragraphs[0].runs[0].bold = True
        sig_table.cell(i, 1).paragraphs[0].runs[0].bold = True


def build_qcto_pm_poe_docx(title: str, syllabus_content: dict, organization_name: str = None,
                            logo_bytes: bytes = None, brand_colors: dict = None,
                            job_id: str = None) -> BytesIO:
    brand_colors = brand_colors or {}
    primary_hex = brand_colors.get("primary", DEFAULT_PRIMARY).lstrip("#") if brand_colors.get("primary") else DEFAULT_PRIMARY
    secondary_hex = brand_colors.get("secondary", DEFAULT_SECONDARY).lstrip("#") if brand_colors.get("secondary") else DEFAULT_SECONDARY
    primary = _hex_to_rgb(primary_hex, DEFAULT_PRIMARY)
    secondary = _hex_to_rgb(secondary_hex, DEFAULT_SECONDARY)
    qualification_title = syllabus_content.get("qualification_title", "") or title

    pm_modules = [m for m in syllabus_content.get("modules", []) if m.get("module_type") == "PM"]

    doc = Document()
    _build_branded_cover(doc, qualification_title, "PM Portfolio of Evidence", organization_name, logo_bytes, primary, primary_hex, secondary)

    _add_poe_cover_details(doc, qualification_title, primary, primary_hex)
    _add_poe_toc(doc, primary, primary_hex, module_count=len(pm_modules) or 1)
    _add_poe_foreword(doc, primary, primary_hex)
    _add_poe_process(doc, primary, primary_hex)
    _add_poe_role_players(doc, primary, primary_hex, secondary)
    _add_poe_competence_and_evidence(doc, primary, primary_hex)
    _add_poe_note_to_learner(doc, primary, primary_hex)
    _add_poe_personal_info(doc, primary, primary_hex)
    _add_poe_preparation_sheet(doc, primary, primary_hex)
    _add_poe_assessment_plan(doc, primary, primary_hex)

    for i, module in enumerate(pm_modules, start=1):
        # Reuses the already-tested per-module intro + scenario/PA-code/IAC-code/
        # Observation-Sheet section functions from the PM Assessment Guide — the learner
        # needs to actually complete this same content, not a re-derived duplicate.
        _add_pag_module_intro(doc, module, i, primary, primary_hex)
        _add_pag_scenario_and_activities(doc, module, i, primary, primary_hex, secondary)

    _add_poe_workplace_diary(doc, primary, primary_hex)
    _add_poe_supervisor_report(doc, primary, primary_hex)
    _add_poe_declaration(doc, qualification_title, primary, primary_hex)

    _add_branded_header_footer(doc, logo_bytes=logo_bytes, qualification_name=qualification_title, organization_name=organization_name, primary_hex=primary_hex)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_qcto_pm_poe_docx_adapter(title, units, organization_name=None, seta=None,
                                     nqf_level=None, logo_bytes=None, brand_colors=None,
                                     job_id=None, **kwargs):
    """Adapter matching the standard DOCUMENT_BUILDERS call signature."""
    syllabus_content = {"modules": units} if isinstance(units, list) else (units or {"modules": []})
    return build_qcto_pm_poe_docx(
        title=title,
        syllabus_content=syllabus_content,
        organization_name=organization_name,
        logo_bytes=logo_bytes,
        brand_colors=brand_colors,
        job_id=job_id,
    )
