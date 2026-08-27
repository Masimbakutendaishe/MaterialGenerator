"""Builds the QCTO ISA (Integrated Summative Assessment) Traceability document — traces
every Knowledge Module topic's Internal Assessment Criteria (IAC, the same field already
extracted as topic assessment_criteria) to any related Practical Module / Workplace Module
content. Linking is done by an AI reasoning pass over the real IAC and PM/WM activity data
(genuine competency overlap, not keyword/title matching) via reason_isa_traceability, with
the original keyword-matching heuristic kept only as an automatic fallback if that call
fails. Topics with no real, extracted IAC get a single clearly-flagged generated
placeholder rather than an empty row, per explicit instruction that IAC 'if not there, are
created'. Topics with no matching PM/WM module are still included, honestly marked as not
directly linked rather than forcing a fabricated match. This table is the specification the
eventual Final Exam is built from."""
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.services.ai_service import reason_isa_traceability
from app.services.document_service import (
    _hex_to_rgb, _build_branded_cover, _add_branded_header_footer, _add_bottom_border,
    _element_text, _element_code, DEFAULT_PRIMARY, DEFAULT_SECONDARY,
)


def _generate_fallback_iac(topic):
    """When a KM topic has no extracted assessment_criteria, generate a single generic
    placeholder criterion rather than leaving the ISA row empty — per explicit instruction
    that IAC 'if not there, are created'. Clearly flagged so it's distinguishable from
    real extracted IAC and can be reviewed by a subject matter expert before use."""
    title = topic.get("title", "this topic")
    return [f"Learners can define, describe, and apply the concepts covered in {title}"]


def _find_linked_modules(km_topic, candidate_modules):
    """Best-effort keyword-overlap linking between a KM topic and PM/WM modules — matches
    on shared significant words (length > 3) in the titles. Returns an empty list (not a
    fabricated link) when no real overlap exists, so unmatched topics stay honestly
    unmatched rather than being forced into a false pairing."""
    km_words = set(w.lower().strip(".,()") for w in km_topic.get("title", "").split() if len(w) > 3)
    linked = []
    for m in candidate_modules:
        m_words = set(w.lower().strip(".,()") for w in m.get("title", "").split() if len(w) > 3)
        if km_words & m_words:
            linked.append(m)
    return linked


def _add_isa_learner_details(doc, primary, primary_hex):
    """Learner and exam-session details, filled in by whoever administers the final exam
    this ISA specifies — placed up front so it's the first thing completed."""
    heading = doc.add_paragraph()
    h_run = heading.add_run("Learner and Exam Details")
    h_run.bold = True
    h_run.font.size = Pt(18)
    h_run.font.color.rgb = primary
    _add_bottom_border(heading, primary_hex)

    table = doc.add_table(rows=7, cols=2)
    table.style = "Table Grid"
    for i, field in enumerate(["Learner Full Name & Surname", "Learner ID Number", "Date of Exam",
                                "Facilitator Name & Number", "Assessor Name & Number",
                                "Moderator Name & Number", "Attempt Number"]):
        table.cell(i, 0).text = field
        table.cell(i, 0).paragraphs[0].runs[0].bold = True
    doc.add_page_break()


def _add_isa_final_signoff(doc, primary, primary_hex):
    """Overall result and sign-off, closing out the learner's exam record against this
    ISA's full traceability table."""
    heading = doc.add_paragraph()
    h_run = heading.add_run("Overall Result and Sign-Off")
    h_run.bold = True
    h_run.font.size = Pt(18)
    h_run.font.color.rgb = primary
    _add_bottom_border(heading, primary_hex)

    doc.add_paragraph(
        "Based on the Criteria Met records above, record the learner's overall result "
        "against this ISA. Any criterion recorded as Not Met should be addressed through "
        "re-assessment before the learner is marked Competent overall."
    )

    judgement_p = doc.add_paragraph()
    judgement_p.add_run("Overall Judgement:  ☐ Competent    ☐ Not Yet Competent").bold = True

    doc.add_paragraph()
    doc.add_paragraph("Assessor Comments:").runs[0].bold = True
    doc.add_paragraph("_" * 100)
    doc.add_paragraph()

    doc.add_paragraph("Moderator Comments:").runs[0].bold = True
    doc.add_paragraph("_" * 100)

    doc.add_paragraph()
    sig_table = doc.add_table(rows=3, cols=2)
    sig_table.style = "Table Grid"
    sig_table.cell(0, 0).text = "Assessor Signature"
    sig_table.cell(0, 1).text = "Date"
    sig_table.cell(1, 0).text = "Moderator Signature"
    sig_table.cell(1, 1).text = "Date"
    sig_table.cell(2, 0).text = "Learner Signature (acknowledging result)"
    sig_table.cell(2, 1).text = "Date"
    for row in sig_table.rows:
        row.cells[0].paragraphs[0].runs[0].bold = True
        row.cells[1].paragraphs[0].runs[0].bold = True


def _add_isa_intro(doc, primary, primary_hex, ai_reasoning_used):
    p = doc.add_paragraph()
    run = p.add_run("Integrated Summative Assessment (ISA) Traceability")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = primary
    _add_bottom_border(p, primary_hex, size="8")

    doc.add_paragraph(
        "This document traces every Knowledge Module topic's Internal Assessment Criteria "
        "(IAC) to any related Practical Module or Workplace Module content. Where a Knowledge "
        "Module topic has no direct Practical or Workplace counterpart, it is still included "
        "here — knowledge, practical, and workplace competence are not always paired one-to-one."
    )
    doc.add_paragraph(
        "This table defines exactly what learners must be able to demonstrate in the Final "
        "Exam. Every Internal Assessment Criterion listed here should be addressed and met "
        "by the Final Exam's questions and practical requirements."
    )
    doc.add_paragraph()
    note = doc.add_paragraph()
    linking_method = (
        "Module links below were determined by an AI reasoning pass over the real IAC and "
        "Practical/Workplace activity content — judging genuine competency overlap, not "
        "matching similar wording."
        if ai_reasoning_used else
        "Module links below were determined by simple keyword overlap between module titles "
        "(the AI reasoning pass was unavailable for this generation) — please verify these "
        "links manually, as this method is less reliable than genuine competency matching."
    )
    note_run = note.add_run(
        "Note: Internal Assessment Criteria marked (generated) were not found explicitly "
        "stated in the source curriculum and have been created as a reasonable placeholder — "
        "these should be reviewed and refined by a subject matter expert before use. "
        f"{linking_method} Module links marked 'Not directly linked' reflect an honest "
        "absence of a matching Practical or Workplace module, not a data error."
    )
    note_run.italic = True
    note_run.font.size = Pt(9)
    doc.add_page_break()


def _add_isa_traceability_table(doc, km_modules, pm_modules, wm_modules, link_lookup, primary, primary_hex, secondary):
    """link_lookup: {topic_code: {"pm": [codes], "wm": [codes], "reasoning": str}} — from
    the AI reasoning pass, or from the keyword-matching fallback if that pass failed."""
    heading = doc.add_paragraph()
    h_run = heading.add_run("Traceability Table")
    h_run.bold = True
    h_run.font.size = Pt(18)
    h_run.font.color.rgb = primary
    _add_bottom_border(heading, primary_hex)

    pm_by_code = {m.get("module_code", ""): m for m in pm_modules}
    wm_by_code = {m.get("module_code", ""): m for m in wm_modules}
    linked_pm_codes_seen = set()
    linked_wm_codes_seen = set()

    def _pm_label(code):
        m = pm_by_code.get(code)
        if not m:
            return code
        pa_items = m.get("performance_assessment") or []
        pa_codes = [_element_code(p) or _element_text(p)[:20] for p in pa_items]
        codes_str = f" ({', '.join(pa_codes)})" if pa_codes else ""
        return f"{code}: {m.get('title', '')}{codes_str}"

    def _wm_label(code):
        m = wm_by_code.get(code)
        if not m:
            return code
        we_items = m.get("work_experience_elements") or []
        we_codes = [_element_code(w) or _element_text(w)[:20] for w in we_items]
        codes_str = f" ({', '.join(we_codes)})" if we_codes else ""
        return f"{code}: {m.get('title', '')}{codes_str}"

    for module in km_modules:
        module_heading = doc.add_paragraph()
        mh_run = module_heading.add_run(f"{module.get('module_code', '')}: {module.get('title', '')}")
        mh_run.bold = True
        mh_run.font.size = Pt(15)
        mh_run.font.color.rgb = secondary

        for topic in module.get("topics", []):
            topic_code = topic.get("topic_code", "")
            topic_p = doc.add_paragraph()
            topic_p.add_run(f"{topic_code}: {topic.get('title', '')}").bold = True

            elements = topic.get("elements") or []
            if elements:
                el_p = doc.add_paragraph()
                el_run = el_p.add_run("Topic Elements: ")
                el_run.bold = True
                el_run.font.size = Pt(9)
                codes_text = ", ".join(
                    f"{_element_code(el)}" if _element_code(el) else _element_text(el)[:30]
                    for el in elements
                )
                el_p.add_run(codes_text).font.size = Pt(9)

            criteria = topic.get("assessment_criteria") or []
            generated = False
            if not criteria:
                criteria = _generate_fallback_iac(topic)
                generated = True

            for criterion in criteria:
                c_p = doc.add_paragraph()
                label = f"{criterion}" + (" (generated)" if generated else "")
                run = c_p.add_run(f"• {label}")
                if generated:
                    run.italic = True

            perf_table = doc.add_table(rows=1, cols=1)
            perf_table.style = "Table Grid"
            perf_cell = perf_table.rows[0].cells[0]
            perf_cell.text = "Criteria Met:  ☐ Yes    ☐ No"
            perf_cell.paragraphs[0].runs[0].bold = True
            comments_p = perf_cell.add_paragraph()
            comments_p.add_run("Assessor Comments: ")
            comments_p.add_run("_" * 60)

            link_info = link_lookup.get(topic_code, {"pm": [], "wm": [], "reasoning": ""})
            linked_pm_codes = link_info.get("pm", [])
            linked_wm_codes = link_info.get("wm", [])
            linked_pm_codes_seen.update(linked_pm_codes)
            linked_wm_codes_seen.update(linked_wm_codes)

            link_table = doc.add_table(rows=1, cols=2)
            link_table.style = "Table Grid"
            link_table.rows[0].cells[0].text = "Linked Practical Module(s) & Elements"
            link_table.rows[0].cells[1].text = "Linked Workplace Module(s) & Elements"
            for c in link_table.rows[0].cells:
                c.paragraphs[0].runs[0].bold = True
            data_row = link_table.add_row().cells
            data_row[0].text = "; ".join(_pm_label(code) for code in linked_pm_codes if code in pm_by_code) or "Not directly linked"
            data_row[1].text = "; ".join(_wm_label(code) for code in linked_wm_codes if code in wm_by_code) or "Not directly linked"

            if link_info.get("reasoning"):
                reasoning_p = doc.add_paragraph()
                reasoning_run = reasoning_p.add_run(link_info["reasoning"])
                reasoning_run.italic = True
                reasoning_run.font.size = Pt(9)
                reasoning_run.font.color.rgb = secondary

            doc.add_paragraph()
        doc.add_page_break()

    unlinked_pm = [m for m in pm_modules if m.get("module_code", "") not in linked_pm_codes_seen]
    unlinked_wm = [m for m in wm_modules if m.get("module_code", "") not in linked_wm_codes_seen]
    if unlinked_pm or unlinked_wm:
        summary_heading = doc.add_paragraph()
        sh_run = summary_heading.add_run("Practical / Workplace Modules Not Linked to Any Knowledge Module Topic")
        sh_run.bold = True
        sh_run.font.size = Pt(14)
        sh_run.font.color.rgb = primary
        doc.add_paragraph(
            "These modules did not match any Knowledge Module topic through the reasoning "
            "pass — listed here so nothing is silently omitted from this document."
        )
        for m in unlinked_pm:
            doc.add_paragraph(f"PM — {m.get('module_code', '')}: {m.get('title', '')}", style="List Bullet")
        for m in unlinked_wm:
            doc.add_paragraph(f"WM — {m.get('module_code', '')}: {m.get('title', '')}", style="List Bullet")


def _build_keyword_link_lookup(km_modules, pm_modules, wm_modules):
    """Fallback lookup, in the same shape as the AI reasoning result, built from the
    original keyword-overlap heuristic — used only if the AI reasoning call fails."""
    lookup = {}
    for module in km_modules:
        for topic in module.get("topics", []):
            code = topic.get("topic_code", "")
            linked_pm = _find_linked_modules(topic, pm_modules)
            linked_wm = _find_linked_modules(topic, wm_modules)
            lookup[code] = {
                "pm": [m.get("module_code", "") for m in linked_pm],
                "wm": [m.get("module_code", "") for m in linked_wm],
                "reasoning": "",
            }
    return lookup


def build_qcto_isa_docx(title: str, syllabus_content: dict, organization_name: str = None,
                         logo_bytes: bytes = None, brand_colors: dict = None,
                         job_id: str = None) -> BytesIO:
    brand_colors = brand_colors or {}
    primary_hex = brand_colors.get("primary", DEFAULT_PRIMARY).lstrip("#") if brand_colors.get("primary") else DEFAULT_PRIMARY
    secondary_hex = brand_colors.get("secondary", DEFAULT_SECONDARY).lstrip("#") if brand_colors.get("secondary") else DEFAULT_SECONDARY
    primary = _hex_to_rgb(primary_hex, DEFAULT_PRIMARY)
    secondary = _hex_to_rgb(secondary_hex, DEFAULT_SECONDARY)
    qualification_title = syllabus_content.get("qualification_title", "") or title

    all_modules = syllabus_content.get("modules", [])
    km_modules = [m for m in all_modules if m.get("module_type") == "KM"]
    pm_modules = [m for m in all_modules if m.get("module_type") == "PM"]
    wm_modules = [m for m in all_modules if m.get("module_type") == "WM"]

    ai_reasoning_used = True
    try:
        link_lookup = reason_isa_traceability(km_modules, pm_modules, wm_modules, job_id=job_id)
    except Exception:
        ai_reasoning_used = False
        link_lookup = _build_keyword_link_lookup(km_modules, pm_modules, wm_modules)

    doc = Document()
    _build_branded_cover(doc, qualification_title, "ISA Traceability Document", organization_name, logo_bytes, primary, primary_hex, secondary)

    _add_isa_learner_details(doc, primary, primary_hex)
    _add_isa_intro(doc, primary, primary_hex, ai_reasoning_used)
    _add_isa_traceability_table(doc, km_modules, pm_modules, wm_modules, link_lookup, primary, primary_hex, secondary)
    _add_isa_final_signoff(doc, primary, primary_hex)

    _add_branded_header_footer(doc, logo_bytes=logo_bytes, qualification_name=qualification_title, organization_name=organization_name, primary_hex=primary_hex)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_qcto_isa_docx_adapter(title, units, organization_name=None, seta=None,
                                  nqf_level=None, logo_bytes=None, brand_colors=None,
                                  job_id=None, **kwargs):
    """Adapter matching the standard DOCUMENT_BUILDERS call signature."""
    syllabus_content = {"modules": units} if isinstance(units, list) else (units or {"modules": []})
    return build_qcto_isa_docx(
        title=title,
        syllabus_content=syllabus_content,
        organization_name=organization_name,
        logo_bytes=logo_bytes,
        brand_colors=brand_colors,
        job_id=job_id,
    )
