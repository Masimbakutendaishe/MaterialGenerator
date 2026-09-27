"""Builds a References document — a list of real, verified source documents (Acts,
regulations, ISO/SANS standards) referenced within a qualification's already-generated
Knowledge Module content, each linked to a genuine PDF found via a real search (never an
AI-invented URL). Scans every KM module's cached generated content, identifies named
documents via extract_referenced_documents, deduplicates by name across modules, then
searches for a verified PDF per unique name. A document no genuine PDF could be found for
is simply omitted from the list, not filled with a guess."""
from io import BytesIO
from docx import Document
from docx.shared import Pt
from app.services.document_service import (
    _hex_to_rgb, _build_branded_cover, _add_branded_header_footer, _add_bottom_border,
    _add_hyperlink, DEFAULT_PRIMARY, DEFAULT_SECONDARY, DEFAULT_ACCENT, _set_default_font,
)


def _extract_block_text(blocks):
    """Pulls all plain text out of a topic/section blocks list, across every block type
    that carries readable text, for scanning by the reference-document extractor."""
    texts = []
    for block in blocks:
        btype = block.get("type")
        if btype == "paragraph":
            texts.append(block.get("text", ""))
        elif btype == "list":
            texts.extend(block.get("items", []))
        elif btype == "example_tip":
            texts.append(block.get("example", ""))
            texts.append(block.get("tip", ""))
        elif btype == "scenario":
            texts.append(block.get("text", ""))
        elif btype == "info_box":
            texts.extend(block.get("items", []))
    return " ".join(t for t in texts if t)


def build_qcto_reference_documents_docx(title: str, syllabus_content: dict, organization_name: str = None,
                                         logo_bytes: bytes = None, brand_colors: dict = None,
                                         accreditation_info: dict = None, job_id: str = None) -> BytesIO:
    from app.services.ai_service import extract_referenced_documents, parallel_map
    from app.services.reference_document_service import find_referenced_document_pdfs

    accreditation_info = dict(accreditation_info or {})
    if syllabus_content.get("qualification_code"):
        accreditation_info.setdefault("qualification_code", syllabus_content.get("qualification_code"))
    brand_colors = brand_colors or {}
    primary_hex = brand_colors.get("primary", DEFAULT_PRIMARY).lstrip("#") if brand_colors.get("primary") else DEFAULT_PRIMARY
    secondary_hex = brand_colors.get("secondary", DEFAULT_SECONDARY).lstrip("#") if brand_colors.get("secondary") else DEFAULT_SECONDARY
    accent_hex = brand_colors.get("accent", DEFAULT_ACCENT).lstrip("#") if brand_colors.get("accent") else DEFAULT_ACCENT
    primary = _hex_to_rgb(primary_hex, DEFAULT_PRIMARY)
    secondary = _hex_to_rgb(secondary_hex, DEFAULT_SECONDARY)
    qualification_title = syllabus_content.get("qualification_title", "") or title

    km_modules = [m for m in syllabus_content.get("modules", []) if m.get("module_type") == "KM"]
    modules_with_content = [m for m in km_modules if m.get("generated_km_content")]

    doc = Document()
    _set_default_font(doc, brand_colors.get("font"))
    _build_branded_cover(doc, qualification_title, "Reference Documents", organization_name, logo_bytes, primary, primary_hex, secondary, accent_hex=accent_hex)

    intro = doc.add_paragraph()
    intro.add_run(
        "This list contains real, source-verified documents referenced within the "
        "Knowledge Module material for this qualification — each link points to a "
        "genuine PDF found via live search, never an AI-generated link. A referenced "
        "document that could not be genuinely verified is not included here; providers "
        "should source it directly if required."
    )
    doc.add_paragraph()

    if not modules_with_content:
        note = doc.add_paragraph()
        note.add_run(
            "No Knowledge Module content has been generated yet for this qualification — "
            "generate the KM Learner Guide first, then regenerate this document to find "
            "its referenced sources."
        ).italic = True
    else:
        # Extract referenced document names across all modules, in parallel
        module_texts = []
        for module in modules_with_content:
            content = module.get("generated_km_content", {})
            all_text = " ".join(_extract_block_text(t.get("blocks", [])) for t in content.get("topics", []))
            module_texts.append(all_text)

        per_module_docs = parallel_map(
            module_texts,
            lambda text: extract_referenced_documents(text, job_id=job_id) if text.strip() else [],
            max_workers=3,
        )

        # Deduplicate by name across all modules
        seen_names = set()
        unique_names = []
        for docs in per_module_docs:
            if not docs:
                continue
            for d in docs:
                name = d.get("name", "").strip()
                if name and name not in seen_names:
                    seen_names.add(name)
                    unique_names.append(name)

        if not unique_names:
            note = doc.add_paragraph()
            note.add_run("No specifically-named reference documents were identified in this qualification's material.").italic = True
        else:
            verified = find_referenced_document_pdfs(unique_names)
            if not verified:
                note = doc.add_paragraph()
                note.add_run(
                    f"{len(unique_names)} referenced document(s) were identified, but a genuine "
                    "PDF could not be verified for any of them via search."
                ).italic = True
            else:
                for ref in verified:
                    p = doc.add_paragraph(style="List Bullet")
                    _add_hyperlink(p, ref["url"], ref["name"])

    _add_branded_header_footer(
        doc, logo_bytes=logo_bytes, qualification_name=qualification_title,
        organization_name=organization_name, primary_hex=primary_hex,
        accreditation_info=accreditation_info, document_label="Reference Documents",
    )

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_qcto_reference_documents_docx_adapter(title, units, organization_name=None, seta=None,
                                                 nqf_level=None, logo_bytes=None, brand_colors=None,
                                                 accreditation_info=None, job_id=None, **kwargs):
    """Adapter matching the standard DOCUMENT_BUILDERS call signature."""
    syllabus_content = {"modules": units} if isinstance(units, list) else (units or {"modules": []})
    return build_qcto_reference_documents_docx(
        title=title, syllabus_content=syllabus_content, organization_name=organization_name,
        logo_bytes=logo_bytes, brand_colors=brand_colors, accreditation_info=accreditation_info,
        job_id=job_id,
    )