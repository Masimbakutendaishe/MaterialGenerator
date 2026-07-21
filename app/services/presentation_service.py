# Builds .pptx from structured content
"""Builds a .pptx presentation from structured syllabus content."""
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


def build_presentation_pptx(title: str, units: list, organization_name: str = None, brand_colors: dict = None) -> BytesIO:
    """Takes syllabus units (list of {"name": ..., "outcomes": [...]}) and produces
    a formatted .pptx in memory. One title slide, then one slide per unit."""
    prs = Presentation()

    primary_color = None
    if brand_colors and brand_colors.get("primary"):
        try:
            hex_color = brand_colors["primary"].lstrip("#")
            primary_color = RGBColor.from_string(hex_color)
        except (ValueError, KeyError):
            primary_color = None  # fall back to default theme color if the hex is malformed

    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = title
    if organization_name:
        slide.placeholders[1].text = organization_name

    if primary_color:
        slide.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = primary_color

    # One slide per unit
    bullet_layout = prs.slide_layouts[1]  # "Title and Content"
    for unit in units:
        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = unit.get("name", "Unit")

        body = slide.placeholders[1].text_frame
        body.clear()
        outcomes = unit.get("outcomes", [])
        if outcomes:
            body.text = outcomes[0]
            for outcome in outcomes[1:]:
                p = body.add_paragraph()
                p.text = outcome

        if primary_color:
            slide.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = primary_color

    buffer = BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer