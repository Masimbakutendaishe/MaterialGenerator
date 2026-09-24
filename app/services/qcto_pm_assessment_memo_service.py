"""Builds a standalone PM Assessment Memorandum — the marking memorandum for the PM
Assessment (Scenario-Based), as its own document. Reuses the exact real questions already
generated for the PM Assessment (via generated_pm_assessment on each module) and the same
generate_model_answers_for_questions call used for KM, so the answers genuinely correspond
to the actual assessment document rather than an independently-generated, possibly
different set of questions."""
from io import BytesIO
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.services.ai_service import generate_model_answers_for_questions
from app.services.document_service import (
    _hex_to_rgb, _build_branded_cover, _add_branded_header_footer, _add_bottom_border,
    DEFAULT_PRIMARY, DEFAULT_SECONDARY, _set_default_font,
)


def build_qcto_pm_assessment_memo_docx(title: str, syllabus_content: dict, organization_name: str = None,
                                        logo_bytes: bytes = None, brand_colors: dict = None,
                                        accreditation_info: dict = None, job_id: str = None) -> BytesIO:
    accreditation_info = dict(accreditation_info or {})
    if syllabus_content.get("qualification_code"):
        accreditation_info.setdefault("qualification_code", syllabus_content.get("qualification_code"))
    brand_colors = brand_colors or {}
    primary_hex = brand_colors.get("primary", DEFAULT_PRIMARY).lstrip("#") if brand_colors.get("primary") else DEFAULT_PRIMARY
    secondary_hex = brand_colors.get("secondary", DEFAULT_SECONDARY).lstrip("#") if brand_colors.get("secondary") else DEFAULT_SECONDARY
    primary = _hex_to_rgb(primary_hex, DEFAULT_PRIMARY)
    secondary = _hex_to_rgb(secondary_hex, DEFAULT_SECONDARY)

    qualification_title = syllabus_content.get("qualification_title", "") or title
    pm_modules = [m for m in syllabus_content.get("modules", []) if m.get("module_type") == "PM"]

    doc = Document()
    _set_default_font(doc, brand_colors.get("font"))
    _build_branded_cover(doc, qualification_title, "PM Assessment Memorandum", organization_name, logo_bytes, primary, primary_hex, secondary)

    note = doc.add_paragraph()
    note_run = note.add_run(
        "Please note that the model answers in this document act as an example of what type "
        "of answers are expected from the learner. Assessors should use discretion and accept "
        "answers in the learner's own words that demonstrate the same understanding."
    )
    note_run.italic = True
    note_run.font.size = Pt(10)
    doc.add_page_break()

    for module in pm_modules:
        real_content = module.get("generated_pm_assessment")

        module_heading = doc.add_paragraph()
        mh_run = module_heading.add_run(f"{module.get('module_code', '')}: {module.get('title', '')}")
        mh_run.bold = True
        mh_run.font.size = Pt(14)
        mh_run.font.color.rgb = primary
        _add_bottom_border(module_heading, primary_hex, size="6")

        if not real_content:
            note_p = doc.add_paragraph()
            note_p.add_run(
                "(No PM Assessment has been generated yet for this module — generate the PM "
                "Assessment document first, then regenerate this memorandum for exact question "
                "correspondence.)"
            ).italic = True
            doc.add_page_break()
            continue

        content = generate_model_answers_for_questions(module, real_content, job_id=job_id)
        for section in content.get("sections", []):
            section_heading = doc.add_paragraph()
            section_heading.add_run(section.get("section_label", "")).bold = True
            for q in section.get("questions", []):
                q_para = doc.add_paragraph()
                q_para.add_run(q.get("question_text", ""))
                marks_run = q_para.add_run(f"  ({q.get('marks', 0)})")
                marks_run.italic = True
                marks_run.font.color.rgb = secondary
                answer_para = doc.add_paragraph()
                answer_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                for point in q.get("model_answer_points", []):
                    run = answer_para.add_run(f"{point}\n")
                    run.italic = True
                    run.font.color.rgb = secondary
                    run.font.size = Pt(11)
                doc.add_paragraph()
        doc.add_page_break()

    _add_branded_header_footer(
        doc, logo_bytes=logo_bytes, qualification_name=qualification_title,
        organization_name=organization_name, primary_hex=primary_hex,
        accreditation_info=accreditation_info, document_label="PM Assessment Memorandum",
    )

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_qcto_pm_assessment_memo_docx_adapter(title, units, organization_name=None, seta=None,
                                                nqf_level=None, logo_bytes=None, brand_colors=None,
                                                accreditation_info=None, job_id=None, **kwargs):
    """Adapter matching the standard DOCUMENT_BUILDERS call signature."""
    syllabus_content = {"modules": units}
    if accreditation_info:
        syllabus_content["qualification_code"] = accreditation_info.get("qualification_code")
        syllabus_content["qualification_title"] = title
    return build_qcto_pm_assessment_memo_docx(
        title=title, syllabus_content=syllabus_content, organization_name=organization_name,
        logo_bytes=logo_bytes, brand_colors=brand_colors, accreditation_info=accreditation_info,
        job_id=job_id,
    )