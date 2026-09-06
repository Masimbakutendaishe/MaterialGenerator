"""Builds the QCTO PM Facilitator Guide — adapted from the KM Facilitator Guide skeleton
(TOC, Introduction, Required Info, Project Assignment, Preparation, Feedback Report,
Facilitation Plan, Marking Memorandum), but reframed throughout around practical,
hands-on skill delivery rather than classroom knowledge transfer: demonstration-based
facilitation methods, practical/workshop accreditation criteria, a Demonstrate → Guided
Practice → Independent Practice → Debrief facilitation cycle per unit (using real PM unit
titles), and a module-level Marking Memorandum matching PM's real data shape (assessment
criteria live at module level, not per-unit, same as the PM Learner Guide's own pattern)."""
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.services.ai_service import generate_facilitator_guide_content, generate_pm_facilitation_steps, parallel_map
from app.services.document_service import (
    _hex_to_rgb, _build_branded_cover, _add_branded_header_footer, _add_bottom_border,
    _add_toc_field, DEFAULT_PRIMARY, DEFAULT_SECONDARY,
)


def _section_heading(doc, text, primary, primary_hex, size=16):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = primary
    p.style = doc.styles["Heading 1"]
    _add_bottom_border(p, primary_hex, size="8")
    return p


def _add_pm_fg_toc(doc, primary, primary_hex, day_count):
    _section_heading(doc, "Table of Contents", primary, primary_hex, size=18)
    _add_toc_field(doc)
    doc.add_page_break()


def _add_pm_fg_introduction(doc, primary, primary_hex, secondary):
    _section_heading(doc, "Facilitator Guide Introduction", primary, primary_hex, size=18)
    _section_heading(doc, "Facilitation Methodology", primary, primary_hex, size=14)

    _section_heading(doc, "Provider Programme Accreditation Criteria", primary, primary_hex, size=12)
    doc.add_paragraph("Physical Requirements:").runs[0].bold = True
    for item in [
        "A practical workshop, simulation area, or workspace equipped for hands-on skill practice",
        "Real or realistic equipment, tools, and materials matching the practical skills being taught",
        "Sufficient space and safety provisions for supervised individual and group practical exercises",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph("Human Resource Requirements:").runs[0].bold = True
    for item in [
        "Facilitators with genuine hands-on practical experience in the skills being taught, not classroom-only expertise",
        "Facilitators who have achieved a nationally accepted standard in the delivery of occupational learning",
        "Facilitators of learning who have achieved a recognised learning standard in assessment practice",
        "Not more than 12 learners per facilitator during hands-on practical sessions, to allow real individual supervision",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph("Legal Requirements: None").runs[0].bold = True
    doc.add_paragraph("Exemptions: None recognised").runs[0].bold = True

    _section_heading(doc, "Methodology Principles", primary, primary_hex, size=12)
    doc.add_paragraph(
        "Practical modules are learned by doing, not by being told. The methodology should "
        "ensure that:"
    )
    for point in [
        "Every skill is demonstrated by the facilitator before learners attempt it themselves.",
        "Learners get guided, supervised practice before being asked to perform independently.",
        "Mistakes made during guided practice are treated as learning opportunities, not failures.",
        "Real or realistic scenarios and equipment are used wherever safely possible.",
        "Each practical session ends with a debrief, so learners can reflect on what worked and what didn't.",
        "Safety is never compromised for the sake of covering more content faster.",
        "The facilitator observes closely enough to catch and correct technique errors before they become habits.",
        "Learners with different starting skill levels are given practice time proportional to their need.",
    ]:
        doc.add_paragraph(point, style="List Bullet")

    _section_heading(doc, "Facilitation Methods", primary, primary_hex, size=12)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Method"
    table.rows[0].cells[1].text = "Description"
    for c in table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True

    methods = [
        ("Demonstration", "The facilitator performs the skill or task first, narrating each step, so learners see exactly what competent performance looks like before attempting it."),
        ("Guided Practice", "Learners attempt the skill with the facilitator actively coaching, correcting technique in real time, and answering questions as they arise."),
        ("Independent Practice", "Learners perform the skill unsupervised (or lightly supervised) to build confidence and fluency, with the facilitator observing from a distance."),
        ("Scenario-Based Exercise", "Learners work through a realistic workplace scenario that requires applying the skill in context, not in isolation."),
        ("Peer Coaching", "Learners observe and give feedback to one another performing the same task, reinforcing the skill from both the doing and the observing side."),
        ("Practical Walkthrough", "The facilitator and learner walk through a completed practical activity together, discussing decisions made and alternatives available."),
    ]
    for name, desc in methods:
        row = table.add_row().cells
        row[0].text, row[1].text = name, desc

    doc.add_paragraph()
    _section_heading(doc, "Assessment", primary, primary_hex, size=12)
    doc.add_paragraph(
        "Formative assessment happens continuously during practical sessions — the facilitator "
        "observes technique, corrects errors, and notes readiness for independent practice. "
        "Summative assessment is conducted through direct observation of a practical "
        "demonstration and/or a Portfolio of Evidence containing real workplace application. "
        "Evidence for practical skills is judged primarily on whether the learner can actually "
        "perform the task to the required standard, not just describe how to perform it."
    )

    _section_heading(doc, "Range Statements", primary, primary_hex, size=12)
    doc.add_paragraph(
        "Also included in the curriculum are the range statements in support of the assessment "
        "criteria, indicating the detailed conditions and context under which the practical "
        "skill must be demonstrated."
    )

    _section_heading(doc, "The Learner Guide", primary, primary_hex, size=12)
    doc.add_paragraph(
        "The PM Learner Guide is included in this material under various practical units, "
        "designed to guide the learner logically through the scope of each practical skill "
        "before they attempt it hands-on."
    )

    _section_heading(doc, "RPL Assessment", primary, primary_hex, size=12)
    doc.add_paragraph(
        "For practical modules, RPL evidence is especially important — assessors should give "
        "particular weight to demonstrated workplace evidence (photos, videos, supervisor "
        "sign-off, completed work samples) over paper certificates alone, since practical "
        "competence is what is actually being assessed. All evidence is verified before a "
        "learner is enrolled for an RPL programme."
    )
    doc.add_page_break()


def _add_pm_fg_programme_info(doc, pm_modules, qualification_title, primary, primary_hex):
    _section_heading(doc, "Required Information About the Programme", primary, primary_hex, size=18)
    doc.add_paragraph(f"Qualification: {qualification_title}")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Module Code", "Title", "NQF Level", "Credits"
    for c in hdr:
        c.paragraphs[0].runs[0].bold = True

    for module in pm_modules:
        row = table.add_row().cells
        row[0].text = module.get("module_code", "")
        row[1].text = module.get("title", "")
        row[2].text = str(module.get("nqf_level", ""))
        row[3].text = str(module.get("credits", ""))
    doc.add_page_break()


def _add_pm_fg_practical_assignment(doc, primary, primary_hex):
    _section_heading(doc, "The Practical Assignment", primary, primary_hex, size=18)

    sub = doc.add_paragraph()
    sub.add_run("Facilitating This Course").bold = True

    doc.add_paragraph("1. Prior to Training").runs[0].bold = True
    for item in [
        "Confirm the practical workspace, equipment, and materials are ready and safe for use.",
        "Learners are instructed to bring a copy of their ID and any required practical attire/PPE to the workshop.",
        "Registration and enrolment of all learners takes place on the first day of the programme.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph("2. During Training").runs[0].bold = True
    for item in [
        "Confirm total attendance with the office and report any problem areas.",
        "Deliver each practical skill via the Demonstrate → Guided Practice → Independent Practice → Debrief cycle.",
        "Supervise all hands-on practice closely enough to correct errors before they become habits.",
        "Ensure any safety requirements specific to the practical activity are enforced throughout.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph("3. After Training").runs[0].bold = True
    for item in [
        "Schedule practical portfolio submission on completion of gathering evidence (photos, videos, supervisor sign-off).",
        "Provide learner guidance and support via email, fax, and telephone.",
        "Conduct integrated assessment of the candidate's practical competence.",
        "Moderate the assessment process.",
        "Endorse achievement and certification with credit allocation.",
        "Follow up with the client.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_page_break()


def _add_pm_fg_preparation(doc, primary, primary_hex):
    _section_heading(doc, "Facilitator Preparation", primary, primary_hex, size=18)

    doc.add_paragraph("Review the material set in the trainer's file, which consists of the following documentation:")
    for item in ["Facilitator Guide", "PM Learner Guide", "Formative Assessment Guide", "Learner POE Guide", "Assessor Guide"]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph()
    doc.add_paragraph("Workshop Resource Requirements:").runs[0].bold = True
    for item in [
        "Trainer's File (to be returned after training)",
        "Course Admin File (to be completed and returned after training)",
        "Practical equipment, tools, and materials matching each skill being taught",
        "Sufficient stock/consumables for each learner to practise independently, not just observe",
        "Personal protective equipment (PPE), where the practical activity requires it",
        "Scenario briefing packs for scenario-based exercises",
        "PM Learner Guide per learner",
        "Formative Assessment Guide per learner",
        "POE Guide per learner",
        "Assessment Guide per candidate",
    ]:
        doc.add_paragraph(item, style="List Number")
    doc.add_page_break()


def _add_pm_fg_feedback_report(doc, primary, primary_hex):
    _section_heading(doc, "Facilitator Feedback Report", primary, primary_hex, size=18)

    doc.add_paragraph("Within 5 days of completing the training, you are required to:")
    for item in [
        "Complete a facilitator feedback report on the programme delivered, noting which practical skills learners found most difficult.",
        "Return all unused material, equipment, the completed course admin file, trainer's file, and signed-out equipment to the office.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph()
    doc.add_paragraph("Preparation Checklist:").runs[0].bold = True
    prep_table = doc.add_table(rows=1, cols=3)
    prep_table.style = "Table Grid"
    prep_table.rows[0].cells[0].text = "Preparation"
    prep_table.rows[0].cells[1].text = "Yes"
    prep_table.rows[0].cells[2].text = "No"
    for c in prep_table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True

    for item in [
        "Qualification Knowledge — I have familiarised myself with the content of the applicable qualification",
        "Curriculum Knowledge — I have familiarised myself with the content of the applicable curriculum",
        "Practical Skill Proficiency — I can personally demonstrate every skill I am about to teach, to the required standard",
        "Equipment Readiness — I have checked all practical equipment and materials are functional and safe",
        "Ability to Respond to Learners' Background — I have studied the learner demographics and prepared accordingly",
        "Safety Preparedness — I know the safety procedures relevant to every practical activity in this session",
    ]:
        row = prep_table.add_row().cells
        row[0].text = item

    doc.add_paragraph()
    doc.add_paragraph("Equipment Check:").runs[0].bold = True
    equip_table = doc.add_table(rows=1, cols=2)
    equip_table.style = "Table Grid"
    equip_table.rows[0].cells[0].text = "Item"
    equip_table.rows[0].cells[1].text = "Checked"
    for c in equip_table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True
    for item in ["PM Learner Guides x 1 per learner", "Assessment Guides x 1 per learner",
                 "Formative Assessment Guides x 1 per learner", "Practical equipment/tools per skill",
                 "Consumable materials, one set per learner", "Personal protective equipment (where required)",
                 "Scenario briefing packs", "First-aid provisions appropriate to the activity"]:
        row = equip_table.add_row().cells
        row[0].text = item

    doc.add_paragraph()
    doc.add_paragraph("Documentation Checklist:").runs[0].bold = True
    doc_table = doc.add_table(rows=1, cols=2)
    doc_table.style = "Table Grid"
    doc_table.rows[0].cells[0].text = "Document"
    doc_table.rows[0].cells[1].text = "Checked"
    for c in doc_table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True
    for item in ["Attendance Register", "Course Evaluation", "Learner Course Evaluation",
                 "Practical Observation Checklists", "Portfolios of Evidence"]:
        row = doc_table.add_row().cells
        row[0].text = item
    doc.add_page_break()


def _add_pm_fg_facilitation_plan(doc, module, day_number, primary, primary_hex, secondary, group_contents=None):
    """Facilitation plan for this module — a lightweight session-flow table for overall
    time planning, followed by genuine, AI-generated, specific step-by-step guidance per
    real practical skill (how to demonstrate it, what to coach for, readiness signs, debrief
    questions) — enough for a first-time facilitator to actually run the session, not a
    template phrase with the skill text substituted in. group_contents is pre-fetched by
    the caller (across the whole document, concurrently) rather than generated here, so
    this function stays purely a renderer and all the slow AI waiting happens once,
    up front, for every module at once."""
    heading = doc.add_paragraph()
    h_run = heading.add_run(f"Facilitation Plan — Day {day_number}: {module.get('title', '')}")
    h_run.bold = True
    h_run.font.size = Pt(16)
    h_run.font.color.rgb = primary
    heading.style = doc.styles["Heading 1"]
    _add_bottom_border(heading, primary_hex)

    doc.add_paragraph(
        "Study the notes in this lesson plan carefully to ensure preparation is done before "
        "the start of class. Personally rehearse every demonstration beforehand — learners "
        "will notice if you're improvising. The step-by-step guidance below assumes no prior "
        "experience facilitating this specific skill."
    )

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Time", "Activity", "Resources"
    for c in hdr:
        c.paragraphs[0].runs[0].bold = True
    for time_str, activity, resources in [
        ("15 min", "Room/Workshop Set Up — ensure venue, equipment, and materials are ready and safe.", ""),
        ("20 min", "Meet, Greet & Seat — learners settle and sign the attendance register.", "Attendance Register, PM Learner Guide"),
        ("15 min", "Overview — welcome, methodology, learner administration and expectations.", "PM Learner Guide"),
    ]:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = time_str, activity, resources

    doc.add_paragraph()

    units = module.get("performance_assessment") or []
    if units and group_contents:
        for facilitation_content in group_contents:
            if facilitation_content is None:
                continue
            for item in facilitation_content.get("items", []):
                skill_heading = doc.add_paragraph()
                skill_heading.add_run(item.get("pa_text", "")).bold = True
                skill_heading.runs[0].font.color.rgb = secondary

                doc.add_paragraph("How to Demonstrate:").runs[0].bold = True
                for i, step in enumerate(item.get("demonstration_steps", []), start=1):
                    doc.add_paragraph(f"{i}. {step}", style="List Number")

                doc.add_paragraph("Coach For (during Guided Practice):").runs[0].bold = True
                for tip in item.get("coaching_tips", []):
                    doc.add_paragraph(tip, style="List Bullet")

                doc.add_paragraph("Signs of Readiness (for Independent Practice):").runs[0].bold = True
                for sign in item.get("readiness_signs", []):
                    doc.add_paragraph(sign, style="List Bullet")

                doc.add_paragraph("Debrief Questions:").runs[0].bold = True
                for q in item.get("debrief_questions", []):
                    doc.add_paragraph(q, style="List Bullet")

                doc.add_paragraph()
    else:
        # No real performance-assessment elements available — say so plainly rather than
        # inventing generic content.
        doc.add_paragraph(
            f"No specific practical skills were extracted for {module.get('title', '')} — "
            f"review the source curriculum for this module before facilitating."
        ).runs[0].italic = True

    end_table = doc.add_table(rows=1, cols=3)
    end_table.style = "Table Grid"
    end_row = end_table.rows[0].cells
    end_row[0].text, end_row[1].text, end_row[2].text = "30 min", "End of Section Parking Bay — take and answer all learner questions.", "PM Learner Guide, Formative Assessment Guide"

    doc.add_page_break()


def build_qcto_pm_facilitator_guide_docx(title: str, syllabus_content: dict, organization_name: str = None,
                                          logo_bytes: bytes = None, brand_colors: dict = None,
                                          accreditation_info: dict = None, job_id: str = None) -> BytesIO:
    accreditation_info = dict(accreditation_info or {})
    if syllabus_content.get("qualification_code"):
        accreditation_info.setdefault("qualification_code", syllabus_content.get("qualification_code"))
    brand_colors = brand_colors or {}
    primary_hex = brand_colors.get("primary", DEFAULT_PRIMARY).lstrip("#") if brand_colors.get("primary") else DEFAULT_PRIMARY
    secondary_hex = brand_colors.get("secondary", DEFAULT_SECONDARY).lstrip("#") if brand_colors.get("secondary") else DEFAULT_SECONDARY
    primary = _hex_to_rgb(primary_hex, DEFAULT_PRIMARY)
    secondary = _hex_to_rgb(secondary_hex, DEFAULT_SECONDARY)
    qualification_title = syllabus_content.get("qualification_title", "") or title

    pm_modules = [m for m in syllabus_content.get("modules", []) if m.get("module_type") == "PM"]

    doc = Document()
    _build_branded_cover(doc, qualification_title, "PM Facilitator Guide", organization_name, logo_bytes, primary, primary_hex, secondary)

    _add_pm_fg_toc(doc, primary, primary_hex, day_count=len(pm_modules) or 1)
    _add_pm_fg_introduction(doc, primary, primary_hex, secondary)
    _add_pm_fg_programme_info(doc, pm_modules, qualification_title, primary, primary_hex)
    _add_pm_fg_practical_assignment(doc, primary, primary_hex)
    _add_pm_fg_preparation(doc, primary, primary_hex)
    _add_pm_fg_feedback_report(doc, primary, primary_hex)

    # Pre-fetch every module's every facilitation-steps group across the WHOLE document
    # concurrently, before any rendering starts — the previous version fetched one group
    # at a time, one module at a time, fully sequentially, which for a curriculum with
    # many modules and many skills per module was the dominant cost in generating this
    # document.
    all_tasks = []
    module_group_counts = []
    for module in pm_modules:
        units = module.get("performance_assessment") or []
        group_size = 3
        groups = [units[i:i + group_size] for i in range(0, len(units), group_size)]
        module_group_counts.append(len(groups))
        for group in groups:
            all_tasks.append((module, group))

    all_results = parallel_map(
        all_tasks,
        lambda t: generate_pm_facilitation_steps(t[0].get("title", ""), t[1], job_id=job_id),
        max_workers=3,
    )

    per_module_results = []
    idx = 0
    for count in module_group_counts:
        per_module_results.append(all_results[idx:idx + count])
        idx += count

    for day_number, module in enumerate(pm_modules, start=1):
        _add_pm_fg_facilitation_plan(doc, module, day_number, primary, primary_hex, secondary, group_contents=per_module_results[day_number - 1])

    # --- Formative Assessment Guide Marking Memorandum — reuses the same generation the
    # KM Facilitator Guide uses, but keyed per MODULE (not per-unit), since PM's real
    # assessment_criteria live at the module level, matching the PM Learner Guide's own
    # reflection-box pattern. ---
    heading = doc.add_paragraph()
    hr = heading.add_run("Formative Assessment Guide Marking Memorandum")
    hr.bold = True
    hr.font.size = Pt(18)
    hr.font.color.rgb = primary
    heading.style = doc.styles["Heading 1"]
    _add_bottom_border(heading, primary_hex)

    note = doc.add_paragraph()
    note_run = note.add_run(
        "Please note that the model answers in this document act as an example of what type "
        "of answers are expected from the learner. Assessors should use discretion and accept "
        "answers in the learner's own words that demonstrate the same understanding. For "
        "practical activities, prioritise direct observation of the skill over written answers."
    )
    note_run.italic = True
    note_run.font.size = Pt(10)

    all_evaluation_items = []

    for module in pm_modules:
        unit_name = module.get("title", "")
        outcomes = module.get("assessment_criteria") or [p.get("text", "") if isinstance(p, dict) else p for p in (module.get("performance_assessment") or [])]

        content = generate_facilitator_guide_content(
            unit_name, outcomes, course_title=qualification_title, job_id=job_id
        )

        section_heading = doc.add_paragraph()
        sh_run = section_heading.add_run(f"{module.get('module_code', '')}: {unit_name}")
        sh_run.bold = True
        sh_run.font.size = Pt(14)
        sh_run.font.color.rgb = secondary

        activity_heading = doc.add_paragraph()
        activity_heading.add_run(content.get("activity_title", "Activity")).bold = True

        for q in content.get("questions", []):
            q_para = doc.add_paragraph()
            q_para.add_run(q.get("question_text", ""))
            marks_run = q_para.add_run(f"  ({q.get('marks', 0)})")
            marks_run.italic = True

            answer_para = doc.add_paragraph()
            answer_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for point in q.get("model_answer_points", []):
                run = answer_para.add_run(f"{point} ✓\n")
                run.italic = True
                run.font.color.rgb = secondary
                run.font.size = Pt(11)

            doc.add_paragraph()

        all_evaluation_items.append({"unit_name": unit_name, "criteria": content.get("evaluation_criteria", [])})

    doc.add_page_break()

    eval_heading = doc.add_paragraph()
    eh_run = eval_heading.add_run("Evaluation Checklist")
    eh_run.bold = True
    eh_run.font.size = Pt(18)
    eh_run.font.color.rgb = primary
    eval_heading.style = doc.styles["Heading 1"]
    _add_bottom_border(eval_heading, primary_hex)

    for label in ["Learner Name:", "Assessor Name:", "Date:"]:
        p = doc.add_paragraph()
        p.add_run(f"{label} " + "_" * 40)

    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Evaluation Criterion"
    hdr[1].text = "Met Requirements"
    hdr[2].text = "Did Not Meet Requirements"
    for cell in hdr:
        cell.paragraphs[0].runs[0].bold = True

    for item in all_evaluation_items:
        for criterion in item["criteria"]:
            row = table.add_row().cells
            row[0].text = criterion

    doc.add_paragraph()
    sig_p = doc.add_paragraph()
    sig_p.add_run("Assessor Signature: " + "_" * 30 + "     Date: " + "_" * 20)

    _add_branded_header_footer(doc, logo_bytes=logo_bytes, qualification_name=qualification_title, organization_name=organization_name, primary_hex=primary_hex, accreditation_info=accreditation_info, document_label="PM Facilitator Guide")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_qcto_pm_facilitator_guide_docx_adapter(title, units, organization_name=None, seta=None,
                                                   nqf_level=None, logo_bytes=None, brand_colors=None,
                                                   accreditation_info=None, job_id=None, **kwargs):
    """Adapter matching the standard DOCUMENT_BUILDERS call signature."""
    syllabus_content = {"modules": units} if isinstance(units, list) else (units or {"modules": []})
    return build_qcto_pm_facilitator_guide_docx(
        title=title,
        syllabus_content=syllabus_content,
        organization_name=organization_name,
        logo_bytes=logo_bytes,
        brand_colors=brand_colors,
        accreditation_info=accreditation_info,
        job_id=job_id,
    )
