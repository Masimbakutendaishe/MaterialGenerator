"""Builds the QCTO KM Facilitator Guide — distinct from the generic Facilitator/Assessor
Guide. Covers: a Table of Contents, a Facilitator Guide Introduction (methodology,
accreditation criteria, delivery methods, assessment approach, RPL), Required Information
About the Programme, the Project Assignment / facilitating-this-course process, Facilitator
Preparation, a Facilitator Feedback Report (with Yes/No preparation checklist, equipment
and documentation checklists), a per-module Facilitation Plan (Time/Activity/Resources),
and a Formative Assessment Guide Marking Memorandum — reusing the same model-answer +
evaluation-checklist generation the existing generic Facilitator Guide uses, matching the
real BCONSULT-style reference structure."""
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.services.ai_service import generate_facilitator_guide_content, generate_model_answers_for_questions, parallel_map
from app.services.document_service import (
    _hex_to_rgb, _build_branded_cover, _add_branded_header_footer, _add_bottom_border,
    _element_text, _add_toc_field, DEFAULT_PRIMARY, DEFAULT_SECONDARY,
)


def _add_km_fg_toc(doc, primary, primary_hex, day_count):
    heading = doc.add_paragraph()
    h_run = heading.add_run("Table of Contents")
    h_run.bold = True
    h_run.font.size = Pt(18)
    h_run.font.color.rgb = primary
    heading.style = doc.styles["Heading 1"]
    _add_bottom_border(heading, primary_hex)

    _add_toc_field(doc)
    doc.add_page_break()


def _add_km_fg_introduction(doc, primary, primary_hex, secondary):
    def _section_heading(text, size=16):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(size)
        run.font.color.rgb = primary
        p.style = doc.styles["Heading 1"]
        _add_bottom_border(p, primary_hex, size="8")
        return p

    _section_heading("Facilitator Guide Introduction", size=18)

    _section_heading("Facilitation Methodology", size=14)

    _section_heading("Provider Programme Accreditation Criteria", size=12)
    doc.add_paragraph("Physical Requirements:").runs[0].bold = True
    for item in [
        "Learning resources aligned to the content of the module",
        "Assessment instruments focused on the internal assessment criteria",
        "Access to training facilities that are well equipped and conducive for effective learning",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph("Human Resource Requirements:").runs[0].bold = True
    for item in [
        "Facilitators of learning with subject matter expertise in the field covered by this module",
        "Facilitators who have achieved a nationally accepted standard in the delivery of occupational learning",
        "Facilitators of learning who have achieved a recognised learning standard in assessment practice",
        "Not more than 20 learners per facilitator",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph("Legal Requirements: None").runs[0].bold = True
    doc.add_paragraph("Exemptions: None recognised").runs[0].bold = True

    _section_heading("Methodology Principles", size=12)
    doc.add_paragraph(
        "The programme is very practical and aims to provide practical tools and skills for "
        "adult learners. The methodology should ensure that:"
    )
    for point in [
        "The learning environment is physically and psychologically comfortable.",
        "Contact training periods are short and varied to avoid boredom.",
        "Learner expectations are articulated, clarified, and managed by the learner and facilitator.",
        "The experience of participants is acknowledged and drawn on in the learning programme.",
        "Facilitation, rather than teaching, is used to allow participants to participate fully.",
        "The facilitator balances new material, debate, and discussion so outcomes are met while all participants are valued.",
        "The learning is problem-oriented, personalised, and accepting of participants' needs for self-direction.",
        "The facilitation process accommodates participants with varying literacy levels.",
    ]:
        doc.add_paragraph(point, style="List Bullet")

    _section_heading("Facilitation Methods", size=12)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Method"
    table.rows[0].cells[1].text = "Description"
    for c in table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True

    methods = [
        ("Structured learning experience", "Participants engage with a complex game or activity representing real-life situations they may encounter in their work."),
        ("Group work", "Participants work on tasks in groups and report their findings back to plenary."),
        ("Lecture", "The facilitator presents a short talk (maximum 10 minutes) to introduce a new subject, provide detail, or wrap up a session."),
        ("Discussion", "A free exchange of ideas or experiences on a particular topic, between facilitator and participants or among participants."),
        ("Brainstorming", "Participants generate ideas on a subject or question, to gather opinions or find out what participants already know."),
        ("Role-play", "Participants act out a scenario, each playing a role, to illustrate how people respond in different situations."),
        ("Case study", "A realistic story or real-life situation where participants apply their knowledge and skills to work through the issues presented."),
    ]
    for name, desc in methods:
        row = table.add_row().cells
        row[0].text, row[1].text = name, desc

    doc.add_paragraph()
    _section_heading("Assessment", size=12)
    doc.add_paragraph(
        "Formative assessment takes place during the learning process, through in-class and "
        "individual exercises that prepare learners for their final assessment. Summative "
        "assessment is conducted at the end of the learning process through a Portfolio of "
        "Evidence. Each outcome has its own set of assessment criteria, describing the evidence "
        "needed to show the learner has demonstrated the outcome correctly."
    )

    _section_heading("Range Statements", size=12)
    doc.add_paragraph(
        "Also included in the curriculum are the range statements in support of the assessment "
        "criteria. The range statements indicate the detailed requirements of the assessment "
        "criteria."
    )

    _section_heading("The Learner Guide", size=12)
    doc.add_paragraph(
        "The Learner Guide is included in this material under various learning units, designed "
        "to guide the learner logically through the material and requirements of the curriculum."
    )

    _section_heading("RPL Assessment", size=12)
    doc.add_paragraph(
        "The assessment of RPL learners is conducted in the same way as for new learners, using "
        "the same assessment pack. RPL applicants must provide proof of previous learning and "
        "subject-related experience prior to assessment — such as certified copies of previous "
        "learning programmes attended. All evidence is assessed and authenticated before a "
        "learner is enrolled for an RPL programme."
    )
    doc.add_page_break()


def _add_km_fg_programme_info(doc, km_modules, qualification_title, primary, primary_hex):
    heading = doc.add_paragraph()
    h_run = heading.add_run("Required Information About the Programme")
    h_run.bold = True
    h_run.font.size = Pt(18)
    h_run.font.color.rgb = primary
    heading.style = doc.styles["Heading 1"]
    _add_bottom_border(heading, primary_hex)

    doc.add_paragraph(f"Qualification: {qualification_title}")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Module Code", "Title", "NQF Level", "Credits"
    for c in hdr:
        c.paragraphs[0].runs[0].bold = True

    for module in km_modules:
        row = table.add_row().cells
        row[0].text = module.get("module_code", "")
        row[1].text = module.get("title", "")
        row[2].text = str(module.get("nqf_level", ""))
        row[3].text = str(module.get("credits", ""))

    doc.add_page_break()


def _add_km_fg_project_assignment(doc, primary, primary_hex):
    heading = doc.add_paragraph()
    h_run = heading.add_run("The Project Assignment")
    h_run.bold = True
    h_run.font.size = Pt(18)
    h_run.font.color.rgb = primary
    heading.style = doc.styles["Heading 1"]
    _add_bottom_border(heading, primary_hex)

    sub = doc.add_paragraph()
    sub.add_run("Facilitating This Course").bold = True

    doc.add_paragraph("1. Prior to Training").runs[0].bold = True
    for item in [
        "Learners are instructed to bring a copy of their ID to the classroom.",
        "Registration and enrolment of all learners takes place on the first day of the programme.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph("2. During Training").runs[0].bold = True
    for item in [
        "Confirm total attendance with the office and report any problem areas.",
        "Deliver the programme via classroom training, self-study where applicable, and workplace application of set tasks and assignments.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph("3. After Training").runs[0].bold = True
    for item in [
        "Schedule portfolio submission on completion of gathering evidence.",
        "Provide learner guidance and support via email, fax, and telephone.",
        "Conduct integrated assessment of the candidate's competence.",
        "Moderate the assessment process.",
        "Endorse achievement and certification with credit allocation.",
        "Follow up with the client.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_page_break()


def _add_km_fg_preparation(doc, primary, primary_hex):
    heading = doc.add_paragraph()
    h_run = heading.add_run("Facilitator Preparation")
    h_run.bold = True
    h_run.font.size = Pt(18)
    h_run.font.color.rgb = primary
    heading.style = doc.styles["Heading 1"]
    _add_bottom_border(heading, primary_hex)

    doc.add_paragraph("Review the material set in the trainer's file, which consists of the following documentation:")
    for item in ["Facilitator Guide", "Learner Guide", "Formative Assessment Guide", "Learner POE Guide", "Assessor Guide"]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph()
    doc.add_paragraph("Workshop Resource Requirements:").runs[0].bold = True
    for item in [
        "Trainer's File (to be returned after training)",
        "Course Admin File (to be completed and returned after training)",
        "Data projector",
        "Suitable screen / surface for projection",
        "Flipchart and flipchart pens",
        "Learner Guide per learner",
        "Formative Assessment Guide per learner",
        "POE Guide per learner",
        "Assessment Guide per candidate",
    ]:
        doc.add_paragraph(item, style="List Number")
    doc.add_page_break()


def _add_km_fg_feedback_report(doc, primary, primary_hex):
    heading = doc.add_paragraph()
    h_run = heading.add_run("Facilitator Feedback Report")
    h_run.bold = True
    h_run.font.size = Pt(18)
    h_run.font.color.rgb = primary
    heading.style = doc.styles["Heading 1"]
    _add_bottom_border(heading, primary_hex)

    doc.add_paragraph("Within 5 days of completing the training, you are required to:")
    for item in [
        "Complete a facilitator feedback report on the programme delivered.",
        "Return all unused material, the completed course admin file, trainer's file, and signed-out equipment to the office.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph()
    doc.add_paragraph("Preparation Checklist:").runs[0].bold = True
    prep_table = doc.add_table(rows=1, cols=3)
    prep_table.style = "Table Grid"
    prep_table.rows[0].cells[0].text = "Preparation"
    prep_table.rows[0].cells[1].text = "Yes"
    prep_table.rows[0].cells[2].text = "No"
    for c in prep_table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True

    prep_items = [
        "Qualification Knowledge — I have familiarised myself with the content of the applicable qualification",
        "Curriculum Knowledge — I have familiarised myself with the content of the applicable curriculum",
        "Content Knowledge — I have sufficient knowledge of the content to facilitate with ease",
        "Application Knowledge — I understand the programme structure and have prepared accordingly",
        "Ability to Respond to Learners' Background — I have studied the learner demographics and prepared accordingly",
        "Enthusiasm & Commitment — I am prepared to create a motivating learning environment",
    ]
    for item in prep_items:
        row = prep_table.add_row().cells
        row[0].text = item

    doc.add_paragraph()
    doc.add_paragraph("Equipment Check:").runs[0].bold = True
    equip_table = doc.add_table(rows=1, cols=2)
    equip_table.style = "Table Grid"
    equip_table.rows[0].cells[0].text = "Item"
    equip_table.rows[0].cells[1].text = "Checked"
    for c in equip_table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True
    for item in ["Learner Guides x 1 per learner", "Assessment Guides x 1 per learner",
                 "Formative Assessment Guides x 1 per learner", "Writing materials & stationery",
                 "White board & pens", "Flip chart paper", "Projector & screen", "Laptop & programme files"]:
        row = equip_table.add_row().cells
        row[0].text = item

    doc.add_paragraph()
    doc.add_paragraph("Documentation Checklist:").runs[0].bold = True
    doc_table = doc.add_table(rows=1, cols=2)
    doc_table.style = "Table Grid"
    doc_table.rows[0].cells[0].text = "Document"
    doc_table.rows[0].cells[1].text = "Checked"
    for c in doc_table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True
    for item in ["Attendance Register", "Course Evaluation", "Learner Course Evaluation", "Portfolios of Evidence"]:
        row = doc_table.add_row().cells
        row[0].text = item
    doc.add_page_break()


def _add_km_fg_facilitation_plan(doc, module, day_number, primary, primary_hex, secondary):
    heading = doc.add_paragraph()
    h_run = heading.add_run(f"Facilitation Plan — Day {day_number}: {module.get('title', '')}")
    h_run.bold = True
    h_run.font.size = Pt(16)
    h_run.font.color.rgb = primary
    heading.style = doc.styles["Heading 1"]
    _add_bottom_border(heading, primary_hex)

    doc.add_paragraph(
        "Study the notes in this lesson plan carefully to ensure preparation is done before "
        "the start of class, and review the Learner Guide to be familiar with the topics that "
        "will be covered."
    )

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Time", "Activity", "Resources"
    for c in hdr:
        c.paragraphs[0].runs[0].bold = True

    rows_data = [
        ("15 min", "Room Set Up — ensure venue and equipment are ready.", ""),
        ("20 min", "Meet, Greet & Seat — learners settle and sign the attendance register.", "Attendance Register, Learner Guide"),
        ("15 min", "Overview — welcome, methodology, learner administration and expectations.", "Learner Guide"),
    ]

    topics = module.get("topics", [])
    for topic in topics:
        rows_data.append(("45 min", f"Knowledge Subject: {topic.get('title', '')}", "Learner Guide"))

    rows_data.append(("30 min", "End of Section Parking Bay — take and answer all learner questions.", "Learner Guide, Formative Assessment Guide"))
    rows_data.append(("45 min", "Lunch Break", ""))

    for time_str, activity, resources in rows_data:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = time_str, activity, resources

    doc.add_page_break()


def build_qcto_km_facilitator_guide_docx(title: str, syllabus_content: dict, organization_name: str = None,
                                          logo_bytes: bytes = None, brand_colors: dict = None,
                                          accreditation_info: dict = None, job_id: str = None) -> BytesIO:
    accreditation_info = dict(accreditation_info or {})
    if syllabus_content.get("qualification_code"):
        accreditation_info.setdefault("qualification_code", syllabus_content.get("qualification_code"))
    brand_colors = brand_colors or {}
    primary_hex = brand_colors.get("primary", DEFAULT_PRIMARY).lstrip("#") if brand_colors.get("primary") else DEFAULT_PRIMARY
    secondary_hex = brand_colors.get("secondary", DEFAULT_SECONDARY).lstrip("#") if brand_colors.get("secondary") else DEFAULT_SECONDARY
    primary = _hex_to_rgb(primary_hex, DEFAULT_PRIMARY)
    secondary = _hex_to_rgb(secondary_hex, DEFAULT_SECONDARY)
    qualification_title = syllabus_content.get("qualification_title", "") or title

    km_modules = [m for m in syllabus_content.get("modules", []) if m.get("module_type") == "KM"]

    doc = Document()
    _build_branded_cover(doc, qualification_title, "KM Facilitator Guide", organization_name, logo_bytes, primary, primary_hex, secondary)

    _add_km_fg_toc(doc, primary, primary_hex, day_count=len(km_modules) or 1)
    _add_km_fg_introduction(doc, primary, primary_hex, secondary)
    _add_km_fg_programme_info(doc, km_modules, qualification_title, primary, primary_hex)
    _add_km_fg_project_assignment(doc, primary, primary_hex)
    _add_km_fg_preparation(doc, primary, primary_hex)
    _add_km_fg_feedback_report(doc, primary, primary_hex)

    for day_number, module in enumerate(km_modules, start=1):
        _add_km_fg_facilitation_plan(doc, module, day_number, primary, primary_hex, secondary)

    # --- Formative Assessment Guide Marking Memorandum (reuses the same generation the
    # existing generic Facilitator Guide uses — same model answers + evaluation checklist
    # logic, just retitled and placed last to match the real reference's TOC order) ---
    heading = doc.add_paragraph()
    hr = heading.add_run("Formative Assessment Guide Marking Memorandum")
    hr.bold = True
    hr.font.size = Pt(18)
    hr.font.color.rgb = primary
    heading.style = doc.styles["Heading 1"]
    _add_bottom_border(heading, primary_hex)

    note = doc.add_paragraph()
    note_run = note.add_run(
        "Please note that the model answers in this document act as an example of what type "
        "of answers are expected from the learner. Assessors should use discretion and accept "
        "answers in the learner's own words that demonstrate the same understanding."
    )
    note_run.italic = True
    note_run.font.size = Pt(10)

    all_evaluation_items = []

    for i, module in enumerate(km_modules, start=1):
        real_content = module.get("generated_formative_assessment")

        module_heading = doc.add_paragraph()
        mh_run = module_heading.add_run(f"{module.get('module_code', '')}: {module.get('title', '')}")
        mh_run.bold = True
        mh_run.font.size = Pt(14)
        mh_run.font.color.rgb = secondary

        if real_content:
            # True correspondence — same real questions as the KM Formative Assessment.
            content = generate_model_answers_for_questions(module, real_content, job_id=job_id)
            for section in content.get("sections", []):
                section_heading = doc.add_paragraph()
                section_heading.add_run(section.get("section_label", "")).bold = True
                for q in section.get("questions", []):
                    q_para = doc.add_paragraph()
                    q_para.add_run(q.get("question_text", ""))
                    marks_run = q_para.add_run(f"  ({q.get('marks', 0)})")
                    marks_run.italic = True
                    answer_para = doc.add_paragraph()
                    answer_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    for point in q.get("model_answer_points", []):
                        run = answer_para.add_run(f"{point} ✓\n")
                        run.italic = True
                        run.font.color.rgb = secondary
                        run.font.size = Pt(11)
                    doc.add_paragraph()
            all_evaluation_items.append({"unit_name": module.get("title", ""), "criteria": content.get("evaluation_criteria", [])})
        else:
            # No real KM Formative Assessment exists yet for this qualification — fall
            # back to independently-generated content per topic (useful standalone
            # content, not guaranteed to match a document that doesn't exist yet) rather
            # than leaving a gap in the guide.
            note_p = doc.add_paragraph()
            note_p.add_run(
                "(Independently generated — regenerate this guide after the KM Formative "
                "Assessment exists for exact question correspondence.)"
            ).italic = True

            topics = module.get("topics", [])
            topic_contents = parallel_map(
                topics,
                lambda t: generate_facilitator_guide_content(
                    t.get("title", ""),
                    t.get("assessment_criteria") or [_element_text(el) for el in t.get("elements") or []],
                    course_title=qualification_title,
                    job_id=job_id,
                ),
                max_workers=3,
            )

            for topic, content in zip(topics, topic_contents):
                unit_name = topic.get("title", "")
                if content is None:
                    continue

                section_heading = doc.add_paragraph()
                sh_run = section_heading.add_run(f"{topic.get('topic_code', '')}: {unit_name}")
                sh_run.bold = True
                sh_run.font.size = Pt(13)

                activity_heading = doc.add_paragraph()
                activity_heading.add_run(content.get("activity_title", "Activity")).bold = True

                for q in content.get("questions", []):
                    q_para = doc.add_paragraph()
                    q_para.add_run(q.get("question_text", ""))
                    marks_run = q_para.add_run(f"  ({q.get('marks', 0)})")
                    marks_run.italic = True
                    answer_para = doc.add_paragraph()
                    answer_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    for point in q.get("model_answer_points", []):
                        run = answer_para.add_run(f"{point} ✓\n")
                        run.italic = True
                        run.font.color.rgb = secondary
                        run.font.size = Pt(11)
                    doc.add_paragraph()

                all_evaluation_items.append({"unit_name": unit_name, "criteria": content.get("evaluation_criteria", [])})

    doc.add_page_break()

    eval_heading = doc.add_paragraph()
    eh_run = eval_heading.add_run("Evaluation Checklist")
    eh_run.bold = True
    eh_run.font.size = Pt(18)
    eh_run.font.color.rgb = primary
    eval_heading.style = doc.styles["Heading 1"]
    _add_bottom_border(eval_heading, primary_hex)

    for label in ["Learner Name:", "Assessor Name:", "Date:"]:
        p = doc.add_paragraph()
        p.add_run(f"{label} " + "_" * 40)

    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Evaluation Criterion"
    hdr[1].text = "Met Requirements"
    hdr[2].text = "Did Not Meet Requirements"
    for cell in hdr:
        cell.paragraphs[0].runs[0].bold = True

    for item in all_evaluation_items:
        for criterion in item["criteria"]:
            row = table.add_row().cells
            row[0].text = criterion

    doc.add_paragraph()
    sig_p = doc.add_paragraph()
    sig_p.add_run("Assessor Signature: " + "_" * 30 + "     Date: " + "_" * 20)

    _add_branded_header_footer(doc, logo_bytes=logo_bytes, qualification_name=qualification_title, organization_name=organization_name, primary_hex=primary_hex, accreditation_info=accreditation_info, document_label="KM Facilitator Guide")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_qcto_km_facilitator_guide_docx_adapter(title, units, organization_name=None, seta=None,
                                                   nqf_level=None, logo_bytes=None, brand_colors=None,
                                                   accreditation_info=None, job_id=None, **kwargs):
    """Adapter matching the standard DOCUMENT_BUILDERS call signature."""
    syllabus_content = {"modules": units} if isinstance(units, list) else (units or {"modules": []})
    return build_qcto_km_facilitator_guide_docx(
        title=title,
        syllabus_content=syllabus_content,
        organization_name=organization_name,
        logo_bytes=logo_bytes,
        brand_colors=brand_colors,
        accreditation_info=accreditation_info,
        job_id=job_id,
    )
