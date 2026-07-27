"""Generalized builder for SETA guide-style documents (Facilitator Guide, Assessment Guide,
Moderator Guide, PoE Guide, etc). Shares content-block rendering with document_service.py."""
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.services.ai_service import generate_guide_section_content
from app.services.document_service import _hex_to_rgb, _add_bottom_border, _render_content_block, _add_page_numbers, DEFAULT_PRIMARY, DEFAULT_SECONDARY, DEFAULT_ACCENT

DOCUMENT_LABELS = {
    "learner_manual": "Learner Manual",
    "facilitator_guide": "Facilitator Guide",
    "formative_assessment": "Formative Assessment",
    "summative_assessment": "Summative Assessment",
    "assessment_guide": "Assessment Guide",
    "moderator_guide": "Moderator Guide",
    "poe_guide": "Portfolio of Evidence Guide",
}


def build_guide_docx(title: str, units: list, organization_name: str = None,
                      seta: str = None, nqf_level: str = None, logo_bytes: bytes = None,
                      brand_colors: dict = None, document_subtype: str = "learner_manual",
                      job_id: str = None) -> BytesIO:
    brand_colors = brand_colors or {}
    primary_hex = brand_colors.get("primary", DEFAULT_PRIMARY).lstrip("#") if brand_colors.get("primary") else DEFAULT_PRIMARY
    secondary_hex = brand_colors.get("secondary", DEFAULT_SECONDARY).lstrip("#") if brand_colors.get("secondary") else DEFAULT_SECONDARY
    primary = _hex_to_rgb(primary_hex, DEFAULT_PRIMARY)
    secondary = _hex_to_rgb(secondary_hex, DEFAULT_SECONDARY)

    label = DOCUMENT_LABELS.get(document_subtype, document_subtype.replace("_", " ").title())

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    from app.services.document_service import _build_branded_cover
    _build_branded_cover(doc, title, label, organization_name, logo_bytes, primary, primary_hex, secondary)

    for unit in units:
        unit_name = unit.get("name", "Unit")
        outcomes = unit.get("outcomes", [])

        heading = doc.add_paragraph()
        heading_run = heading.add_run(unit_name)
        heading_run.bold = True
        heading_run.font.size = Pt(18)
        heading_run.font.color.rgb = primary
        _add_bottom_border(heading, primary_hex)

        content = generate_guide_section_content(document_subtype, unit_name, outcomes, seta=seta, nqf_level=nqf_level, job_id=job_id)

        if content.get("intro"):
            intro_p = doc.add_paragraph()
            intro_run = intro_p.add_run(content["intro"])
            intro_run.italic = True

        for section in content.get("sections", []):
            if section.get("heading"):
                sec_heading = doc.add_paragraph()
                sec_run = sec_heading.add_run(section["heading"])
                sec_run.bold = True
                sec_run.underline = True
                sec_run.font.size = Pt(14)
                sec_run.font.color.rgb = secondary

            for block in section.get("blocks", []):
                _render_content_block(doc, block, primary_hex, secondary)

        key_points = content.get("key_points", [])
        if key_points:
            kp_heading = doc.add_paragraph()
            kp_run = kp_heading.add_run("KEY POINTS")
            kp_run.bold = True
            kp_run.underline = True
            for point in key_points:
                doc.add_paragraph(point, style="List Bullet")

        doc.add_page_break()

    _add_page_numbers(doc)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer