"""Builds the QCTO KM Assessment Guide — an assessor-facing policy/process document,
distinct from the KM Formative Assessment (the actual questions the learner answers) and
the KM Facilitator Guide's Marking Memorandum (the answer key). This covers: the 4-step
assessment process, role-player responsibilities, assessment principles, evidence types
(VARCS), assessment methods, learner preparation and special-needs checklists, assessment
details/logistics tied to the real qualification and modules, a Summative Assessment
section that cross-references the real KM Formative Assessment's topics, and the
appeal/re-assessment policy — matching the real BCONSULT-style Assessor Guide reference.
Document names referenced throughout ("KM Learner Guide", "KM Facilitator Guide", "KM
Formative Assessment", "KM Video Guide") match the real document titles this system
generates, not generic placeholders, so the guide actually links up to what the learner
is reading elsewhere."""
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.services.document_service import (
    _hex_to_rgb, _build_branded_cover, _add_branded_header_footer, _add_bottom_border,
    DEFAULT_PRIMARY, DEFAULT_SECONDARY,
)


def _section_heading(doc, text, primary, primary_hex, size=18):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = primary
    _add_bottom_border(p, primary_hex, size="8")
    return p


def _add_ag_cover_details(doc, qualification_title, primary, primary_hex):
    """Candidate/assessor fillable info table right after the cover — Learner Name, ID,
    Assessor, Date, matching the reference's cover-page table."""
    doc.add_paragraph(qualification_title).runs[0].bold = True

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    fields = ["Learner Full Name & Surname", "ID Number", "Assessor", "Date"]
    for i, field in enumerate(fields):
        table.cell(i, 0).text = field
        table.cell(i, 0).paragraphs[0].runs[0].bold = True
    doc.add_page_break()


def _add_ag_toc(doc, primary, primary_hex):
    _section_heading(doc, "Table of Contents", primary, primary_hex)
    for entry in [
        "Foreword to the Assessor",
        "Assessment Process",
        "Assessment Steps",
        "The Assessment Process Role-players",
        "Principles of Assessment",
        "Assessing Evidence",
        "Assessment Methods",
        "Assessment Review",
        "Learning Programme and Assessment Strategy Notes",
        "Learner Preparation for Assessment Checklist",
        "Special Need Record",
        "Assessment Details",
        "Assessment Logistics",
        "Summative Assessment Marking Memo",
        "Learner Portfolio of Evidence Guide",
        "Appeal Policy and Procedure",
        "Appeal Form",
        "Re-Assessment Policy",
    ]:
        doc.add_paragraph(entry, style="List Bullet")
    doc.add_page_break()


def _add_ag_foreword(doc, primary, primary_hex):
    _section_heading(doc, "Foreword to the Assessor", primary, primary_hex)
    doc.add_paragraph(
        "The purpose of this guide is to provide the Assessor with guidelines on the "
        "assessment of the Portfolio of Evidence submitted by the candidate (learner) for "
        "assessment purposes."
    )
    doc.add_paragraph(
        "Assessment in Outcomes-Based Education is not only focused on what learners can do, "
        "but intends to develop learners holistically. Learners are also required to "
        "demonstrate certain life skills, which enhance their learning and are transferable "
        "to their private lives."
    )
    doc.add_paragraph(
        "Formative assessment refers to assessment that takes place during the process of "
        "learning and teaching. Summative assessment is assessment for making a judgement "
        "about achievement, carried out when a learner is ready to be assessed at the end of "
        "a programme of learning. Results initially collected for formative assessment can be "
        "used for summative assessment, avoiding repetition."
    )
    doc.add_page_break()


def _add_ag_process(doc, primary, primary_hex):
    _section_heading(doc, "Assessment Process", primary, primary_hex)
    doc.add_paragraph(
        "The assessment process follows four steps, in order: plan and prepare for the "
        "assessment; conduct and record the assessment; provide feedback on the assessment; "
        "and review and report on the assessment."
    )
    for i, step in enumerate([
        "Plan & prepare for assessment",
        "Conduct & record the assessment",
        "Provide feedback on the assessment",
        "Review and report on the assessment",
    ], start=1):
        doc.add_paragraph(f"Step {i}: {step}", style="List Number")
    doc.add_page_break()


def _add_ag_steps(doc, qualification_title, primary, primary_hex, secondary):
    """The four assessment steps, each with Assessor needs / Learner needs / Documents —
    referencing the real KM document names this system generates."""
    _section_heading(doc, "Assessment Steps", primary, primary_hex)

    km_learner_guide = "KM Learner Guide"
    km_facilitator_guide = "KM Facilitator Guide"
    km_formative_assessment = "KM Formative Assessment"

    steps = [
        (
            "Step 1: Plan and Prepare for the Assessment",
            [
                "Understand and review all requirements of the assessment in terms of evidence required to prove competence.",
                "Identify logistical arrangements that have to be made, such as the venue.",
                f"Familiarise themselves with the {km_facilitator_guide} and its assessment instruments and tools.",
                "Ensure familiarity with the related Assessment, Moderation, RPL, and Appeals policies.",
            ],
            [
                "Be informed of the assessment requirements, roles, responsibilities, and how evidence is to be collected.",
                f"Be guided in preparing for assessment using the {km_learner_guide} and {km_formative_assessment}.",
                "Be given the contact details of the facilitator, assessor, and other support persons.",
                "Have any assessment process related questions answered.",
            ],
            f"{km_learner_guide}, {km_formative_assessment}, {km_facilitator_guide}",
        ),
        (
            "Step 2: Conduct and Record the Assessment",
            [
                "Conduct the assessment in an appropriate, non-threatening manner, applying the assessment principles.",
                f"Review and assess the evidence submitted by the learner, referring to the {km_facilitator_guide}'s Marking Memorandum for guidelines and model answers.",
                "Make a judgement about the evidence against the assessment criteria, using the principles of good evidence as a guideline.",
                "Record the assessment process, findings, and decisions in the required format.",
            ],
            [
                f"Submit their completed {km_formative_assessment} and any Portfolio of Evidence items for review.",
            ],
            f"{km_formative_assessment}, {km_facilitator_guide} Marking Memorandum",
        ),
        (
            "Step 3: Provide Assessment Feedback to the Learner",
            [
                "Provide the learner with feedback in both a positive and constructive manner.",
                "Advise the learner of any outstanding or required evidence.",
                "Record all communication with the learner.",
            ],
            [
                "Confirm receipt, understanding, and acceptance of the feedback.",
            ],
            "Assessment Feedback document",
        ),
        (
            "Step 4: Review and Report the Assessment",
            [
                "Review the assessment process and report on it.",
                "Record the outcome of the assessment in the relevant quality management system.",
                "Manage any learner appeals against the assessment outcome, according to the Appeal Policy and Procedure.",
            ],
            [
                "Review the assessment process by completing a review questionnaire.",
            ],
            "Assessor and Moderator Review of the Assessment",
        ),
    ]

    for title, assessor_items, learner_items, documents in steps:
        step_heading = doc.add_paragraph()
        step_heading.add_run(title).bold = True
        step_heading.runs[0].font.size = Pt(14)
        step_heading.runs[0].font.color.rgb = secondary

        doc.add_paragraph("Assessor needs to:").runs[0].bold = True
        for item in assessor_items:
            doc.add_paragraph(item, style="List Bullet")

        doc.add_paragraph("Learner needs to:").runs[0].bold = True
        for item in learner_items:
            doc.add_paragraph(item, style="List Bullet")

        doc.add_paragraph(f"Documents: {documents}").runs[0].italic = True
        doc.add_paragraph()

    doc.add_page_break()


def _add_ag_role_players(doc, primary, primary_hex, secondary):
    _section_heading(doc, "The Assessment Process Role-Players", primary, primary_hex)

    learner_heading = doc.add_paragraph()
    learner_heading.add_run("Learner").bold = True
    learner_heading.runs[0].font.color.rgb = secondary
    doc.add_paragraph(
        "Learners participate in the facilitated classroom training by taking part in "
        "formative assessment class activities and exercises found in the KM Formative "
        "Assessment. The learner needs to attend the learning session, participate in group "
        "work, research and prepare for participation, and complete all assignments, "
        "activities, and portfolio requirements."
    )
    doc.add_paragraph(
        "Assessments are meant to be clear and transparent, so learners should know: the "
        "kinds of assessment activities they will perform, the standard and level of "
        "performance expected, the type and amount of evidence to be collected, and their "
        "responsibility regarding the collection of evidence."
    )

    facilitator_heading = doc.add_paragraph()
    facilitator_heading.add_run("Facilitator").bold = True
    facilitator_heading.runs[0].font.color.rgb = secondary
    doc.add_paragraph(
        "It is the role of the facilitator to facilitate the theoretical classroom training "
        "and skills practice sessions, using the KM Learner Guide and KM Video Guide. The "
        "facilitator is also responsible for being available for assessment-related questions, "
        "guiding learners through the formative assessment activities, and handling learning "
        "programme administration."
    )
    doc.add_page_break()


def _add_ag_principles(doc, primary, primary_hex):
    _section_heading(doc, "Principles of Assessment", primary, primary_hex)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Principle"
    table.rows[0].cells[1].text = "Description"
    for c in table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True

    principles = [
        ("Valid", "The assessment focuses on the requirements laid down in the standard — it is fit for purpose."),
        ("Authentic", "The assessor is satisfied that the work being assessed is attributable to the person being assessed."),
        ("Current", "The evidence must demonstrate that the person's competence is current."),
        ("Sufficient", "The evidence establishes that all criteria are met and performance can be repeated consistently."),
        ("Fair", "The method chosen must be fair and must not present barriers unrelated to the evidence."),
        ("Systematic", "The assessment must be planned and recorded to ensure it is fair."),
        ("Appropriate", "The method of assessment is suited to the performance being assessed."),
        ("Manageable", "The methods used must be accessible, cost-effective, and not interfere with learning."),
        ("Open", "The candidate contributes to the planning and collecting of evidence and understands the process."),
        ("Consistent", "The same assessor makes the same judgement under similar circumstances."),
        ("Integrated", "Evidence collection is integrated into the work or learning process where feasible."),
        ("Direct", "The activities in the assessment mirror the conditions of actual performance as closely as possible."),
        ("Time Efficient", "The assessment does not interfere with the candidate's normal daily activities."),
    ]
    for name, desc in principles:
        row = table.add_row().cells
        row[0].text, row[1].text = name, desc
    doc.add_page_break()


def _add_ag_evidence(doc, primary, primary_hex):
    _section_heading(doc, "Assessing Evidence", primary, primary_hex)
    doc.add_paragraph(
        "A criterion-based assessment can only be performed using evidence generated by the "
        "learner. Types of evidence include:"
    )
    for label, desc in [
        ("Direct", "Evidence collected directly by the assessor, such as observing the learner perform a task."),
        ("Indirect", "Evidence the learner has collected, signed off as authentic, and submitted for assessment."),
        ("Historic", "Evidence of competence assessed by someone else, such as a prior certificate of competence."),
    ]:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(desc)

    doc.add_paragraph()
    doc.add_paragraph("Evidence must meet the VARCS criteria:").runs[0].bold = True
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Criterion"
    table.rows[0].cells[1].text = "Description"
    for c in table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True
    for letter, desc in [
        ("Valid", "The unit standard or qualification being assessed must require the evidence submitted."),
        ("Authentic", "Evidence submitted must be the learner's own work."),
        ("Relevant", "Evidence must relate directly to the assessment criteria."),
        ("Current", "The evidence must demonstrate that competence is current."),
        ("Sufficient", "The evidence must satisfy all assessment criteria — partial evidence is not sufficient."),
    ]:
        row = table.add_row().cells
        row[0].text, row[1].text = letter, desc
    doc.add_page_break()


def _add_ag_methods(doc, primary, primary_hex):
    _section_heading(doc, "Assessment Methods", primary, primary_hex)
    doc.add_paragraph(
        "Various methods of assessment can be used. The method chosen must be appropriate "
        "for the specific assessment criteria and the context of the assessment."
    )
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Method"
    table.rows[0].cells[1].text = "Description"
    for c in table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True

    methods = [
        ("Observation", "The assessor observes a learner carrying out an activity as a normal part of workplace responsibilities."),
        ("Simulations", "Reproduces the essential characteristics of a real working experience, especially where the real situation would be hazardous."),
        ("Role Playing", "Enables the learner to demonstrate interpersonal and communication skills through enacted scenarios."),
        ("Case Studies", "Evaluates a learner's ability to understand a scenario and respond to relevant questions."),
        ("Demonstrating & Questioning", "The assessor observes a structured practical activity and questions the learner on their actions."),
        ("Pen and Paper Tests", "Written tests measuring the extent of a learner's factual knowledge."),
        ("Oral Tests", "Used to test speed, accuracy, or recall, and useful where literacy is not a competence requirement."),
        ("Projects", "Largely unsupervised, longer-term activities carried out in the workplace or learning environment."),
        ("Portfolios", "A collection of evidence relating to the work being assessed."),
        ("Self-Assessment", "The learner checks their own understanding against the assessment criteria throughout the learning process."),
        ("Computer-Based Assessment", "The learner interacts with a computer programme designed to assess knowledge and skill."),
    ]
    for name, desc in methods:
        row = table.add_row().cells
        row[0].text, row[1].text = name, desc
    doc.add_page_break()


def _add_ag_review_and_strategy(doc, primary, primary_hex):
    _section_heading(doc, "Assessment Review", primary, primary_hex)
    doc.add_paragraph(
        "Reviewing the assessment process could involve: consulting the learner for feedback "
        "about the assessment and how it could be improved; reviewing the process with other "
        "assessors; making appropriate changes; and using the assessment results to evaluate "
        "the learning programme and strategies used."
    )

    _section_heading(doc, "Learning Programme and Assessment Strategy Notes", primary, primary_hex, size=16)
    doc.add_paragraph(
        "As the assessor, ensure familiarity with the learning programme and assessment "
        "strategy for the qualification you are assessing. Review the curriculum document to "
        "understand how assessment activities align with the specific outcomes and assessment "
        "criteria, the assessment methods associated with each activity, and the assessment, "
        "moderation, RPL, and appeals policies and procedures relevant to the training provider."
    )
    doc.add_page_break()


def _add_ag_learner_prep_checklist(doc, primary, primary_hex):
    _section_heading(doc, "Learner Preparation for Assessment Checklist", primary, primary_hex)

    info_table = doc.add_table(rows=4, cols=2)
    info_table.style = "Table Grid"
    for i, field in enumerate(["Learner Name", "Assessor Name", "Date", "Venue"]):
        info_table.cell(i, 0).text = field
        info_table.cell(i, 0).paragraphs[0].runs[0].bold = True

    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "The assessor explained/gave me:"
    table.rows[0].cells[1].text = "Yes"
    table.rows[0].cells[2].text = "No"
    for c in table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True

    for item in [
        "How my assessment is linked to the NQF",
        "The unit standards/curriculum I was to be assessed against",
        "The assessment methods for this assessment",
        "The feedback process and appeals procedure",
        "Any barriers that could affect the fairness of the assessment",
        "Any special assessment requirements (language, disabilities)",
        "A copy of the curriculum/qualification content to be assessed against",
        "The assessment instruments to be used",
        "An opportunity to point out any special needs",
    ]:
        row = table.add_row().cells
        row[0].text = item
    doc.add_page_break()


def _add_ag_special_needs(doc, primary, primary_hex):
    _section_heading(doc, "Special Need Record", primary, primary_hex)

    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Name of Learner"
    table.cell(0, 1).text = "Name of Assessor"
    table.cell(1, 0).text = "Signature of Learner"
    table.cell(1, 1).text = "Signature of Assessor"
    for row in table.rows:
        for cell in row.cells:
            cell.paragraphs[0].runs[0].bold = True

    doc.add_paragraph()
    for label in [
        "List the special needs of the candidate:",
        "Indicate how provision will be made to accommodate these special needs:",
        "List possible barriers that can have a negative influence on the assessment process:",
        "Indicate how these barriers will be accommodated:",
    ]:
        doc.add_paragraph(label).runs[0].bold = True
        doc.add_paragraph("_" * 100)
        doc.add_paragraph()
    doc.add_page_break()


def _add_ag_assessment_details(doc, qualification_title, km_modules, primary, primary_hex):
    """Data-driven assessment details — Programme name and real module codes as the
    'Unit Standards'/modules being assessed, not generic placeholders."""
    _section_heading(doc, "Assessment Details", primary, primary_hex)

    module_codes = ", ".join(m.get("module_code", "") for m in km_modules) or "—"

    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Programme"
    table.cell(0, 1).text = qualification_title
    table.cell(1, 0).text = "Knowledge Modules Assessed"
    table.cell(1, 1).text = module_codes
    table.cell(0, 0).paragraphs[0].runs[0].bold = True
    table.cell(1, 0).paragraphs[0].runs[0].bold = True

    doc.add_paragraph()
    doc.add_paragraph("Attached Documents Checklist:").runs[0].bold = True
    for item in [
        "Learner's Personal Information Document",
        "Assessment Preparation Document",
        "Declaration of Authenticity Document",
        "KM Formative Assessment (completed)",
    ]:
        doc.add_paragraph(f"{item} — Completed: ☐", style="List Bullet")

    doc.add_paragraph()
    sig_table = doc.add_table(rows=6, cols=2)
    sig_table.style = "Table Grid"
    for i, field in enumerate(["Assessor Name", "Assessor Number", "Assessment Date",
                                "Moderator Name", "Moderator Number", "Moderation Date"]):
        sig_table.cell(i, 0).text = field
        sig_table.cell(i, 0).paragraphs[0].runs[0].bold = True

    doc.add_paragraph()
    final_p = doc.add_paragraph()
    final_p.add_run("Final Judgement:  ☐ Competent    ☐ Not Yet Competent").bold = True
    doc.add_page_break()


def _add_ag_logistics(doc, primary, primary_hex):
    _section_heading(doc, "Assessment Logistics", primary, primary_hex)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Activity", "Date", "Evidence", "Place"
    for c in hdr:
        c.paragraphs[0].runs[0].bold = True

    for activity in ["Training Completed", "Self-Assessment", "Assessment Contract",
                      "Initial Meeting", "Knowledge Assessment", "Feedback", "Moderation", "Re-Assessment"]:
        row = table.add_row().cells
        row[0].text = activity
    doc.add_page_break()


def _add_ag_marking_memo(doc, km_modules, primary, primary_hex):
    """Cross-references the real KM Formative Assessment's topics rather than generic
    'Knowledge Questions' — the assessor is pointed at the actual topics this learner's
    KM Formative Assessment covers."""
    _section_heading(doc, "Summative Assessment Marking Memo", primary, primary_hex)
    doc.add_paragraph(
        "The learner needs to individually complete the summative assessment activities in "
        "the KM Formative Assessment. Assessor, use the KM Facilitator Guide's Marking "
        "Memorandum as a guide to assess the following knowledge areas in the learner's "
        "submission:"
    )
    for module in km_modules:
        for topic in module.get("topics", []):
            doc.add_paragraph(f"{topic.get('topic_code', '')}: {topic.get('title', '')}", style="List Bullet")

    doc.add_paragraph()
    doc.add_paragraph("Assessor Notes:").runs[0].bold = True
    doc.add_paragraph(
        "Use the KM Facilitator Guide's Marking Memorandum to mark the knowledge questions. "
        "Complete the Assessment Details and Assessment Logistics sections of this guide to "
        "record your findings."
    )
    doc.add_page_break()


def _add_ag_poe_guide(doc, primary, primary_hex):
    _section_heading(doc, "Learner Portfolio of Evidence Guide", primary, primary_hex)

    doc.add_paragraph("Knowledge Questions").runs[0].bold = True
    doc.add_paragraph(
        "The learner answers the knowledge questions in the KM Formative Assessment, based "
        "on the theory covered in the KM Learner Guide. Assessor Notes: use the KM "
        "Facilitator Guide's Marking Memorandum to mark the knowledge questions, and record "
        "your findings in the Assessment Details section of this guide."
    )

    doc.add_paragraph()
    doc.add_paragraph("Practical Activities").runs[0].bold = True
    doc.add_paragraph(
        "The learner individually completes practical activities to show their ability to "
        "apply their knowledge and skills. Assessor Notes: refer to any observation or "
        "product evaluation evidence submitted, and record your findings in the Assessment "
        "Details section of this guide."
    )
    doc.add_page_break()


def _add_ag_appeal_and_reassessment(doc, primary, primary_hex):
    _section_heading(doc, "Appeal Policy and Procedure", primary, primary_hex)
    doc.add_paragraph(
        "A learner who is not satisfied with an assessment decision has the right to appeal. "
        "The responsibility for implementing the requirements of this policy rests with the "
        "assessor and the training provider's administrator. Appeals must be lodged in "
        "writing within a reasonable period of the assessment outcome being communicated, "
        "using the Appeal Form below."
    )
    doc.add_page_break()

    _section_heading(doc, "Appeal Form", primary, primary_hex)
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    for i, field in enumerate(["Learner Name", "Assessment Date", "Assessor Name", "Reason for Appeal", "Learner Signature & Date"]):
        table.cell(i, 0).text = field
        table.cell(i, 0).paragraphs[0].runs[0].bold = True
    doc.add_page_break()

    _section_heading(doc, "Re-Assessment Policy", primary, primary_hex)
    doc.add_paragraph(
        "When learners undergo re-assessment, they must be given feedback so they can "
        "concentrate on areas of weakness and only be re-assessed on Not Yet Competent "
        "criteria. Re-assessment should take place in the same situation or context and "
        "under the same conditions; the same method and assessment instrument may be used, "
        "but the task and materials should be changed."
    )


def build_qcto_km_assessment_guide_docx(title: str, syllabus_content: dict, organization_name: str = None,
                                         logo_bytes: bytes = None, brand_colors: dict = None,
                                         job_id: str = None) -> BytesIO:
    brand_colors = brand_colors or {}
    primary_hex = brand_colors.get("primary", DEFAULT_PRIMARY).lstrip("#") if brand_colors.get("primary") else DEFAULT_PRIMARY
    secondary_hex = brand_colors.get("secondary", DEFAULT_SECONDARY).lstrip("#") if brand_colors.get("secondary") else DEFAULT_SECONDARY
    primary = _hex_to_rgb(primary_hex, DEFAULT_PRIMARY)
    secondary = _hex_to_rgb(secondary_hex, DEFAULT_SECONDARY)
    qualification_title = syllabus_content.get("qualification_title", "") or title

    km_modules = [m for m in syllabus_content.get("modules", []) if m.get("module_type") == "KM"]

    doc = Document()
    _build_branded_cover(doc, qualification_title, "KM Assessment Guide", organization_name, logo_bytes, primary, primary_hex, secondary)

    _add_ag_cover_details(doc, qualification_title, primary, primary_hex)
    _add_ag_toc(doc, primary, primary_hex)
    _add_ag_foreword(doc, primary, primary_hex)
    _add_ag_process(doc, primary, primary_hex)
    _add_ag_steps(doc, qualification_title, primary, primary_hex, secondary)
    _add_ag_role_players(doc, primary, primary_hex, secondary)
    _add_ag_principles(doc, primary, primary_hex)
    _add_ag_evidence(doc, primary, primary_hex)
    _add_ag_methods(doc, primary, primary_hex)
    _add_ag_review_and_strategy(doc, primary, primary_hex)
    _add_ag_learner_prep_checklist(doc, primary, primary_hex)
    _add_ag_special_needs(doc, primary, primary_hex)
    _add_ag_assessment_details(doc, qualification_title, km_modules, primary, primary_hex)
    _add_ag_logistics(doc, primary, primary_hex)
    _add_ag_marking_memo(doc, km_modules, primary, primary_hex)
    _add_ag_poe_guide(doc, primary, primary_hex)
    _add_ag_appeal_and_reassessment(doc, primary, primary_hex)

    _add_branded_header_footer(doc, logo_bytes=logo_bytes, qualification_name=qualification_title, organization_name=organization_name, primary_hex=primary_hex)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_qcto_km_assessment_guide_docx_adapter(title, units, organization_name=None, seta=None,
                                                  nqf_level=None, logo_bytes=None, brand_colors=None,
                                                  job_id=None, **kwargs):
    """Adapter matching the standard DOCUMENT_BUILDERS call signature."""
    syllabus_content = {"modules": units} if isinstance(units, list) else (units or {"modules": []})
    return build_qcto_km_assessment_guide_docx(
        title=title,
        syllabus_content=syllabus_content,
        organization_name=organization_name,
        logo_bytes=logo_bytes,
        brand_colors=brand_colors,
        job_id=job_id,
    )
