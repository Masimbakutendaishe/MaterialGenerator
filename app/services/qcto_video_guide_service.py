"""Builds the QCTO Video Resource Guide — a preface page plus, for every module across
KM/PM/WM, an AI-curated list of topics each paired with a real clickable YouTube search
link, a description, and learning outcomes."""
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from urllib.parse import quote_plus
from app.services.ai_service import generate_qcto_video_guide_content
from app.services.image_service import fetch_stock_photo
from app.services.document_service import (
        _hex_to_rgb, _build_branded_cover, _add_branded_header_footer, _add_bottom_border, _add_hyperlink,
    DEFAULT_PRIMARY, DEFAULT_SECONDARY,
)


def build_qcto_video_guide_docx(title: str, syllabus_content: dict, organization_name: str = None,
                                 logo_bytes: bytes = None, brand_colors: dict = None,
                                 job_id: str = None) -> BytesIO:
    brand_colors = brand_colors or {}
    primary_hex = brand_colors.get("primary", DEFAULT_PRIMARY).lstrip("#") if brand_colors.get("primary") else DEFAULT_PRIMARY
    secondary_hex = brand_colors.get("secondary", DEFAULT_SECONDARY).lstrip("#") if brand_colors.get("secondary") else DEFAULT_SECONDARY
    primary = _hex_to_rgb(primary_hex, DEFAULT_PRIMARY)
    secondary = _hex_to_rgb(secondary_hex, DEFAULT_SECONDARY)
    qualification_title = syllabus_content.get("qualification_title", "") or title

    doc = Document()
    _build_branded_cover(doc, qualification_title, "KM Video Guide", organization_name, logo_bytes, primary, primary_hex, secondary)

    preface_heading = doc.add_paragraph()
    preface_run = preface_heading.add_run("Preface")
    preface_run.bold = True
    preface_run.font.size = Pt(16)
    preface_run.font.color.rgb = primary
    _add_bottom_border(preface_heading, primary_hex)

    preface_text = doc.add_paragraph(
        "This Video Resource Guide is designed to serve as your go-to reference for high-quality "
        "video tutorials and demonstrations that support your learning and ongoing skills development. "
        "Each entry below links to a targeted search for videos relevant to a specific topic covered "
        "in this qualification. All you need is an internet connection and a device to view them."
    )
    preface_text.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_page_break()

    all_modules = [m for m in syllabus_content.get("modules", []) if m.get("module_type") in ("KM", "PM", "WM")]

    for module in all_modules:
        module_heading = doc.add_paragraph()
        mh_run = module_heading.add_run(f"Module: {module.get('title', '')}")
        mh_run.bold = True
        mh_run.font.size = Pt(16)
        mh_run.font.color.rgb = primary
        _add_bottom_border(module_heading, primary_hex)

        content = generate_qcto_video_guide_content(module, job_id=job_id)

        for entry in content.get("entries", []):
            topic_p = doc.add_paragraph()
            topic_run = topic_p.add_run(entry.get("topic_name", ""))
            topic_run.bold = True
            topic_run.font.size = Pt(13)
            topic_run.font.color.rgb = secondary

            desc_label = doc.add_paragraph()
            desc_label.add_run("Description:").bold = True
            doc.add_paragraph(entry.get("description", ""))

            search_query = entry.get("search_query", "")

            # A relevant photo for visual context — NOT a real video thumbnail (we can't verify
            # real video URLs, so we never fabricate one); this just makes the entry feel like
            # a real resource card rather than plain text.
            try:
                photo_bytes = fetch_stock_photo(search_query)
                if photo_bytes:
                    from io import BytesIO as _BytesIO
                    doc.add_picture(_BytesIO(photo_bytes), width=Inches(2.5))
            except Exception:
                pass

            link_label = doc.add_paragraph()
            link_label.add_run("▶ Watch related videos:").bold = True

            link_p = doc.add_paragraph()
            search_url = f"https://www.youtube.com/results?search_query={quote_plus(search_query)}"
            _add_hyperlink(link_p, search_url, f"Search YouTube: {search_query}")

            outcomes_label = doc.add_paragraph()
            outcomes_label.add_run("Learning Outcomes:").bold = True
            for outcome in entry.get("learning_outcomes", []):
                doc.add_paragraph(outcome, style="List Bullet")

            doc.add_paragraph()

        doc.add_page_break()

    _add_branded_header_footer(doc, logo_bytes=logo_bytes, qualification_name=title, organization_name=organization_name, primary_hex=primary_hex)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_qcto_video_guide_docx_adapter(title, units, organization_name=None, seta=None,
                                         nqf_level=None, logo_bytes=None, brand_colors=None,
                                         job_id=None, **kwargs):
    """Adapter matching the standard DOCUMENT_BUILDERS call signature."""
    syllabus_content = {"modules": units} if isinstance(units, list) else (units or {"modules": []})
    return build_qcto_video_guide_docx(
        title=title,
        syllabus_content=syllabus_content,
        organization_name=organization_name,
        logo_bytes=logo_bytes,
        brand_colors=brand_colors,
        job_id=job_id,
    )
