"""Builds a standalone KM Portfolio of Evidence Memorandum — the marking memorandum for
the KM POE's Formative Assessment questions, as its own document. Reuses the exact real
questions already generated for the KM POE (via generated_km_poe_questions on each topic)
so the answers genuinely correspond to the actual POE document rather than an
independently-generated, possibly different set of questions."""
from io import BytesIO
from docx import Document
from docx.shared import Pt
from app.services.document_service import (
    _hex_to_rgb, _build_branded_cover, _add_branded_header_footer, _add_bottom_border,
    DEFAULT_PRIMARY, DEFAULT_SECONDARY, DEFAULT_ACCENT, _set_default_font,
)


def _render_formative_answer(doc, q_number, q, secondary):
    q_para = doc.add_paragraph()
    q_run = q_para.add_run(f"{q_number}. {q.get('question_text', '')}")
    q_run.bold = True
    marks_run = q_para.add_run(f"  ({q.get('marks', 0)} marks)")
    marks_run.italic = True
    marks_run.font.color.rgb = secondary

    q_type = q.get("type", "open")
    if q_type == "multiple_choice":
        options = q.get("options", {})
        correct = q.get("correct", "")
        for letter in ["A", "B", "C", "D"]:
            if letter in options:
                label = f"{letter}. {options[letter]}" + ("  ✓ CORRECT" if letter == correct else "")
                p = doc.add_paragraph(label, style="List Bullet")
                if letter == correct:
                    p.runs[0].bold = True
                    p.runs[0].font.color.rgb = secondary
    else:
        answer_para = doc.add_paragraph()
        for point in q.get("model_answer_points", []):
            run = answer_para.add_run(f"{point}\n")
            run.italic = True
            run.font.color.rgb = secondary
            run.font.size = Pt(11)
    doc.add_paragraph()


def build_qcto_km_poe_memo_docx(title: str, syllabus_content: dict, organization_name: str = None,
                                 logo_bytes: bytes = None, brand_colors: dict = None,
                                 accreditation_info: dict = None, job_id: str = None) -> BytesIO:
    accreditation_info = dict(accreditation_info or {})
    if syllabus_content.get("qualification_code"):
        accreditation_info.setdefault("qualification_code", syllabus_content.get("qualification_code"))
    brand_colors = brand_colors or {}
    primary_hex = brand_colors.get("primary", DEFAULT_PRIMARY).lstrip("#") if brand_colors.get("primary") else DEFAULT_PRIMARY
    secondary_hex = brand_colors.get("secondary", DEFAULT_SECONDARY).lstrip("#") if brand_colors.get("secondary") else DEFAULT_SECONDARY
    accent_hex = brand_colors.get("accent", DEFAULT_ACCENT).lstrip("#") if brand_colors.get("accent") else DEFAULT_ACCENT
    primary = _hex_to_rgb(primary_hex, DEFAULT_PRIMARY)
    secondary = _hex_to_rgb(secondary_hex, DEFAULT_SECONDARY)
    qualification_title = syllabus_content.get("qualification_title", "") or title

    km_modules = [m for m in syllabus_content.get("modules", []) if m.get("module_type") == "KM"]

    doc = Document()
    _set_default_font(doc, brand_colors.get("font"))
    _build_branded_cover(doc, qualification_title, "KM POE Memorandum", organization_name, logo_bytes, primary, primary_hex, secondary, accent_hex=accent_hex)

    note = doc.add_paragraph()
    note_run = note.add_run(
        "Please note that the model answers in this document act as an example of what type "
        "of answers are expected from the learner. Assessors should use discretion and accept "
        "answers in the learner's own words that demonstrate the same understanding."
    )
    note_run.italic = True
    note_run.font.size = Pt(10)
    doc.add_page_break()

    activity_number = 1
    for module in km_modules:
        module_heading = doc.add_paragraph()
        mh_run = module_heading.add_run(f"{module.get('module_code', '')}: {module.get('title', '')}")
        mh_run.bold = True
        mh_run.font.size = Pt(14)
        mh_run.font.color.rgb = primary
        _add_bottom_border(module_heading, primary_hex, size="6")

        for topic in module.get("topics", []):
            if not topic.get("assessment_criteria"):
                continue
            content = topic.get("generated_km_poe_questions")

            act_heading = doc.add_paragraph()
            act_heading.add_run(f"Activity {activity_number} — {topic.get('title', '')}").bold = True
            act_heading.runs[0].font.color.rgb = secondary
            act_heading.runs[0].font.size = Pt(13)

            if not content:
                note_p = doc.add_paragraph()
                note_p.add_run(
                    "(No KM POE has been generated yet for this topic — generate the KM POE "
                    "document first, then regenerate this memorandum for exact question "
                    "correspondence.)"
                ).italic = True
            else:
                for q_num, q in enumerate(content.get("questions", []), start=1):
                    _render_formative_answer(doc, q_num, q, secondary)

            doc.add_paragraph()
            activity_number += 1
        doc.add_page_break()

    _add_branded_header_footer(
        doc, logo_bytes=logo_bytes, qualification_name=qualification_title,
        organization_name=organization_name, primary_hex=primary_hex,
        accreditation_info=accreditation_info, document_label="KM POE Memorandum",
    )

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_qcto_km_poe_memo_docx_adapter(title, units, organization_name=None, seta=None,
                                         nqf_level=None, logo_bytes=None, brand_colors=None,
                                         accreditation_info=None, job_id=None, **kwargs):
    """Adapter matching the standard DOCUMENT_BUILDERS call signature."""
    syllabus_content = {"modules": units}
    if accreditation_info:
        syllabus_content["qualification_code"] = accreditation_info.get("qualification_code")
        syllabus_content["qualification_title"] = title
    return build_qcto_km_poe_memo_docx(
        title=title, syllabus_content=syllabus_content, organization_name=organization_name,
        logo_bytes=logo_bytes, brand_colors=brand_colors, accreditation_info=accreditation_info,
        job_id=job_id,
    )