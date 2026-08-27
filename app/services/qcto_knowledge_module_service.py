"""Builds the QCTO Knowledge Module document — cover page, front matter (welcome,
methodology, learner admin, learner expectations, icons legend), TOC, per-module
intro/purpose, a sub-modules/units table, detailed per-topic content with example/tip
callouts, per-topic outcome/percentage dividers and reflection boxes, and back matter
(annexures + bibliography) — matching the real structural pattern of accredited QCTO
training material."""
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from app.services.ai_service import generate_qcto_knowledge_module_content
from app.services.document_service import (
    _hex_to_rgb, _build_branded_cover, _add_signature_block, _add_page_numbers,
    _add_branded_header_footer, _render_content_block, _add_bottom_border, _shade_paragraph,
    _element_text, _element_code, DEFAULT_PRIMARY, DEFAULT_SECONDARY, DEFAULT_ACCENT,
)


def _add_km_front_matter(doc, primary, primary_hex, secondary):
    """Static front-matter pages every accredited KM learner guide carries: welcome,
    programme methodology, types of activities, learner administration/support,
    learner expectations (fillable), and the icons legend."""

    def _section_heading(text, size=16):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(size)
        run.font.color.rgb = primary
        _add_bottom_border(p, primary_hex, size="8")
        return p

    _section_heading("Welcome to the Programme")
    doc.add_paragraph(
        "Follow along in this guide as your facilitator takes you through the material. "
        "Make notes and sketches that will help you understand and remember what you have "
        "learnt. Take notes and share information with your colleagues — important and "
        "relevant information and skills are transferred by sharing."
    )

    _section_heading("Programme Methodology")
    doc.add_paragraph(
        "This programme is delivered through facilitator presentations, readings, individual "
        "activities, group discussions, and skill-application exercises. Know what you want "
        "to get out of the programme from the beginning, and start applying your new skills "
        "immediately."
    )

    activities_heading = doc.add_paragraph()
    activities_heading.add_run("Types of Activities You Can Expect").bold = True

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(["Type of Activity", "Description", "Purpose"]):
        hdr[i].text = text
        hdr[i].paragraphs[0].runs[0].bold = True
        hdr[i].paragraphs[0].runs[0].font.size = Pt(10)

    rows_data = [
        ("Knowledge Training", "Completed on your own.",
         "Tests your understanding and ability to apply the information."),
        ("Skills Application Activities", "Completed in the workplace.",
         "Requires you to apply the knowledge and skills gained in the workplace."),
        ("Applied Knowledge", "Collecting information and samples of documents from the workplace.",
         "Ensures you get the opportunity to learn from experts in the industry."),
    ]
    for a, b, c in rows_data:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = a, b, c
        for cell in row:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()

    _section_heading("Learner Administration")
    doc.add_paragraph(
        "You are required to sign the Attendance Register every day you attend training "
        "sessions facilitated by a facilitator. On completion, you will be asked to complete "
        "a Learning Programme Evaluation Form to help improve our service and material."
    )

    _section_heading("Learner Support")
    doc.add_paragraph(
        "The responsibility for learning rests with you — be proactive, ask questions, and "
        "seek assistance from your facilitator when required. You are responsible for the "
        "safekeeping of your completed Formative Assessment Workbook and Workplace Guide."
    )

    doc.add_page_break()

    _section_heading("Learner Expectations")
    doc.add_paragraph(
        "Please prepare the following information. You will be asked to introduce yourself "
        "to the facilitator and your fellow learners."
    )

    for label in ["Your Name", "The Organisation You Represent", "Your Position in the Organisation",
                  "What do you hope to achieve by attending this programme?"]:
        box_table = doc.add_table(rows=2, cols=1)
        box_table.style = "Table Grid"
        label_cell = box_table.rows[0].cells[0]
        label_run = label_cell.paragraphs[0].add_run(label)
        label_run.bold = True
        blank_cell = box_table.rows[1].cells[0]
        blank_cell.paragraphs[0].add_run(" ")
        blank_cell.paragraphs[0].paragraph_format.space_before = Pt(18)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    doc.add_page_break()

    _section_heading("Icons Used in This Programme")
    doc.add_paragraph(
        "As you progress through the guide, you will encounter the following icons, used "
        "to indicate how you are expected to participate at that point."
    )

    icon_table = doc.add_table(rows=1, cols=2)
    icon_table.style = "Table Grid"
    ihdr = icon_table.rows[0].cells
    ihdr[0].text, ihdr[1].text = "Icon", "Description"
    for c in ihdr:
        c.paragraphs[0].runs[0].bold = True

    icon_rows = [
        ("📖", "Introduction to a topic"),
        ("📝", "Space available to make your own notes"),
        ("💡", "Important facts / hints / tips"),
        ("👥", "You will be required to perform an activity under the guidance of your facilitator"),
        ("🧠", "Examples"),
    ]
    for icon, desc in icon_rows:
        row = icon_table.add_row().cells
        row[0].text, row[1].text = icon, desc

    doc.add_page_break()


def _add_topic_outcome_box(doc, module, topic, primary, primary_hex):
    """Section-divider box preceding a topic's detailed content: outcome statement +
    percentage weighting, matching the boxed 'KM-0X-KTXX' pattern in accredited material."""
    table = doc.add_table(rows=2, cols=1)
    table.style = "Table Grid"

    title_cell = table.rows[0].cells[0]
    title_p = title_cell.paragraphs[0]
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(f"{topic.get('topic_code', '')}: {topic.get('title', '')}")
    title_run.bold = True
    title_run.font.size = Pt(13)
    title_run.font.color.rgb = primary

    pct_cell = table.rows[1].cells[0]
    pct_p = pct_cell.paragraphs[0]
    pct_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    weight = topic.get("weight")
    pct_run = pct_p.add_run(f"Percentage: {weight}" if weight else "Percentage: —")
    pct_run.bold = True
    pct_run.font.size = Pt(10)
    _shade_paragraph(pct_p, primary_hex)
    for run in pct_p.runs:
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    doc.add_paragraph()


def _add_reflection_box(doc, module, topic, primary_hex, secondary_hex):
    """'Write your own notes to remember' reflection box closing out a topic, numbering
    the topic's real assessment criteria as reflection prompts (padded to 5 lines) and
    tagging the box with an IAC-style code so it maps back to the curriculum."""
    prompts = topic.get("assessment_criteria") or []

    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]

    heading_run = cell.paragraphs[0].add_run("Write your own notes to remember:")
    heading_run.bold = True
    heading_run.font.size = Pt(11)
    heading_run.font.color.rgb = _hex_to_rgb(primary_hex, DEFAULT_PRIMARY)

    for i, prompt in enumerate(prompts[:5], start=1):
        p = cell.add_paragraph()
        p.add_run(f"{i}. {prompt}")
        p.paragraph_format.space_after = Pt(6)
    for i in range(len(prompts[:5]) + 1, 6):
        p = cell.add_paragraph()
        p.add_run(f"{i}. " + "_" * 60)
        p.paragraph_format.space_after = Pt(6)

    code_p = doc.add_paragraph()
    code_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    code_run = code_p.add_run(
        f"IAC-{module.get('module_code', '')}-{topic.get('topic_code', '')}: "
        f"{topic.get('title', '')} can be explained and demonstrated"
    )
    code_run.italic = True
    code_run.bold = True
    code_run.font.size = Pt(10)
    code_run.font.color.rgb = _hex_to_rgb(secondary_hex, DEFAULT_SECONDARY)
    doc.add_paragraph()


def _add_km_back_matter(doc, primary, primary_hex):
    """Static back-matter annexures every accredited KM learner guide carries: growth
    action plan, new-words glossary, training evaluation, and bibliography."""

    def _section_heading(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = primary
        _add_bottom_border(p, primary_hex, size="8")
        return p

    doc.add_page_break()
    _section_heading("Annexure 1: Growth Action Plan")
    doc.add_paragraph(
        "List, in order of priority, the areas in which you need to improve in order to "
        "become competent."
    )
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, text in enumerate(["Actions to be Taken", "Resources", "Completion Date", "Evidence"]):
        table.rows[0].cells[i].text = text
        table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for _ in range(5):
        table.add_row()

    doc.add_page_break()
    _section_heading("Annexure 2: Words That Are New to Me")
    table2 = doc.add_table(rows=1, cols=2)
    table2.style = "Table Grid"
    table2.rows[0].cells[0].text = "Term"
    table2.rows[0].cells[1].text = "Description"
    for c in table2.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True
    for _ in range(8):
        table2.add_row()

    doc.add_page_break()
    _section_heading("Annexure 3: Training Evaluation")
    doc.add_paragraph("Rate the following from 1 (Poor) to 5 (Excellent).")
    table3 = doc.add_table(rows=1, cols=2)
    table3.style = "Table Grid"
    table3.rows[0].cells[0].text = "Question"
    table3.rows[0].cells[1].text = "Rating (1-5)"
    for c in table3.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True
    for q in [
        "Did the training relate to your job (skills, knowledge)?",
        "To what extent will your performance improve as a result of this training?",
        "Would you recommend this course to others?",
        "Did this training meet your desired needs?",
        "Was the training material user-friendly and easy to understand?",
    ]:
        row = table3.add_row().cells
        row[0].text = q

    doc.add_page_break()
    _section_heading("Bibliography")
    doc.add_paragraph(
        "Acknowledgements & References. This material was compiled using the QCTO "
        "curriculum document and standard project/occupational management references."
    )


def build_qcto_knowledge_module_docx(title: str, syllabus_content: dict, organization_name: str = None,
                                      logo_bytes: bytes = None, brand_colors: dict = None,
                                      job_id: str = None) -> BytesIO:
    brand_colors = brand_colors or {}
    primary_hex = brand_colors.get("primary", DEFAULT_PRIMARY).lstrip("#") if brand_colors.get("primary") else DEFAULT_PRIMARY
    secondary_hex = brand_colors.get("secondary", DEFAULT_SECONDARY).lstrip("#") if brand_colors.get("secondary") else DEFAULT_SECONDARY
    accent_hex = brand_colors.get("accent", DEFAULT_ACCENT).lstrip("#") if brand_colors.get("accent") else DEFAULT_ACCENT
    primary = _hex_to_rgb(primary_hex, DEFAULT_PRIMARY)
    secondary = _hex_to_rgb(secondary_hex, DEFAULT_SECONDARY)

    qualification_code = syllabus_content.get("qualification_code", "")
    qualification_title = syllabus_content.get("qualification_title", "") or title

    doc = Document()
    _build_branded_cover(
        doc, qualification_title, "KM Learner Guide", organization_name, logo_bytes, primary, primary_hex, secondary,
        accent_hex=accent_hex, qualification_code=qualification_code,
    )

    _add_km_front_matter(doc, primary, primary_hex, secondary)

    # Learner/Facilitator/Date lines on their own page after the front matter
    for label in ["Learner Name:", "Facilitator Name:", "Date of Submission:"]:
        p = doc.add_paragraph()
        p.add_run(f"{label} " + "_" * 40)
        p.paragraph_format.space_after = Pt(14)
    doc.add_page_break()

    km_modules = [m for m in syllabus_content.get("modules", []) if m.get("module_type") == "KM"]

    # Table of Contents (simple styled list — real page numbers aren't feasible without
    # Word's native TOC field mechanics, so this lists modules/topics for navigation reference)
    toc_heading = doc.add_paragraph()
    toc_run = toc_heading.add_run("Table of Contents")
    toc_run.bold = True
    toc_run.font.size = Pt(18)
    toc_run.font.color.rgb = primary
    _add_bottom_border(toc_heading, primary_hex)

    for m_index, module in enumerate(km_modules, start=1):
        mod_p = doc.add_paragraph()
        mod_p.add_run(f"Module {m_index}: {module.get('title', '')}").bold = True
        for topic in module.get("topics", []):
            topic_p = doc.add_paragraph()
            topic_p.paragraph_format.left_indent = Inches(0.3)
            topic_p.add_run(f"{topic.get('topic_code', '')} — {topic.get('title', '')}")

    doc.add_page_break()

    # Per-module content
    for m_index, module in enumerate(km_modules, start=1):
        module_heading = doc.add_paragraph()
        mh_run = module_heading.add_run(f"Module {m_index}: {module.get('title', '')}")
        mh_run.bold = True
        mh_run.font.size = Pt(20)
        mh_run.font.color.rgb = primary
        _add_bottom_border(module_heading, primary_hex)

        meta_p = doc.add_paragraph()
        meta_run = meta_p.add_run(f"{module.get('module_code', '')}  |  NQF Level {module.get('nqf_level', '')}  |  {module.get('credits', '')} Credits")
        meta_run.italic = True
        meta_run.font.size = Pt(10)

        content = generate_qcto_knowledge_module_content(module, job_id=job_id)

        intro_heading = doc.add_paragraph()
        intro_heading.add_run("Introduction").bold = True
        intro_p = doc.add_paragraph(content.get("module_intro", ""))
        intro_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        purpose_heading = doc.add_paragraph()
        purpose_heading.add_run("Purpose").bold = True
        purpose_p = doc.add_paragraph(content.get("module_purpose", ""))
        purpose_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        # Sub-modules and Units table
        subheading = doc.add_paragraph()
        subheading.add_run("Sub-modules and Units").bold = True

        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Learning Unit Code"
        hdr[1].text = "Learning Unit Title"
        for cell in hdr:
            cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].runs[0].font.size = Pt(10)

        for topic in module.get("topics", []):
            row = table.add_row().cells
            row[0].text = topic.get("topic_code", "")
            row[1].text = topic.get("title", "")
            for cell in row:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)

        doc.add_paragraph()

        # Per-topic detailed content — matched by topic_code (not position), since the AI
        # response isn't guaranteed to preserve source order/count exactly.
        source_topics_by_code = {t.get("topic_code"): t for t in module.get("topics", [])}
        for topic_content in content.get("topics", []):
            topic_source = source_topics_by_code.get(topic_content.get("topic_code"), {})
            _add_topic_outcome_box(doc, module, topic_source, primary, primary_hex)

            topic_heading = doc.add_paragraph()
            th_run = topic_heading.add_run(f"{topic_content.get('topic_code', '')} — {topic_content.get('topic_title', '')}")
            th_run.bold = True
            th_run.font.size = Pt(15)
            th_run.font.color.rgb = secondary

            topic_elements = topic_source.get("elements") or []
            if topic_elements:
                ref_p = doc.add_paragraph()
                ref_run = ref_p.add_run("Topic Elements Covered:")
                ref_run.bold = True
                ref_run.italic = True
                ref_run.font.size = Pt(10)
                ref_run.font.color.rgb = secondary
                for el in topic_elements:
                    code = _element_code(el)
                    text = _element_text(el)
                    label = f"{code}: {text}" if code else text
                    item_p = doc.add_paragraph(label, style="List Bullet")
                    for run in item_p.runs:
                        run.font.size = Pt(10)
                        run.italic = True

            for block in topic_content.get("blocks", []):
            
                _render_content_block(doc, block, primary_hex, secondary, accent_hex=accent_hex)

            _add_reflection_box(doc, module, topic_source, primary_hex, secondary_hex)

        doc.add_page_break()

    _add_km_back_matter(doc, primary, primary_hex)

    _add_signature_block(doc)
    _add_branded_header_footer(
        doc, logo_bytes=logo_bytes, qualification_name=qualification_title,
        organization_name=organization_name, primary_hex=primary_hex,
    )

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_qcto_knowledge_module_docx_adapter(title, units, organization_name=None, seta=None,
                                              nqf_level=None, logo_bytes=None, brand_colors=None,
                                              job_id=None, **kwargs):
    """Adapter matching the standard DOCUMENT_BUILDERS call signature (title, units, ...)
    used by generate_package_document_task, translating it to this builder's actual
    signature (syllabus_content instead of units, no seta/nqf_level)."""
    syllabus_content = {"modules": units} if isinstance(units, list) else (units or {"modules": []})
    return build_qcto_knowledge_module_docx(
        title=title,
        syllabus_content=syllabus_content,
        organization_name=organization_name,
        logo_bytes=logo_bytes,
        brand_colors=brand_colors,
        job_id=job_id,
    )