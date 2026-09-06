"""Builds the QCTO KM Learner Workbook — an active-learning companion the learner works
through during training, distinct from the KM Learner Guide (explanatory content), KM
Formative Assessment (the graded test), and KM POE (the formal submission package). For
each real KM topic: a key-terms fill-in exercise, self-check confidence ratings built
directly from the topic's real assessment_criteria (not fabricated questions), free
note-taking space, and a single reflective takeaway line."""
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.services.document_service import (
    _hex_to_rgb, _build_branded_cover, _add_branded_header_footer, _add_bottom_border,
    _element_text, _element_code, DEFAULT_PRIMARY, DEFAULT_SECONDARY,
)


def _section_heading(doc, text, primary, primary_hex, size=16):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = primary
    _add_bottom_border(p, primary_hex, size="8")
    return p


def _add_workbook_intro(doc, primary, primary_hex):
    _section_heading(doc, "How to Use This Workbook", primary, primary_hex, size=18)
    doc.add_paragraph(
        "This workbook is yours to write in. Use it alongside the KM Learner Guide during "
        "your training sessions — it is not marked or submitted for assessment, so be honest "
        "with yourself about what you do and don't yet understand."
    )
    doc.add_paragraph("For each topic, you'll find:").runs[0].bold = True
    for item in [
        "Key Terms — jot down any word or concept that's new to you, in your own words.",
        "Self-Check — rate honestly how confident you feel about each part of the topic.",
        "My Notes — space to write whatever helps you remember what was covered.",
        "One Thing I Want to Remember — a single takeaway from the session.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_page_break()


def _add_module_divider(doc, module, primary, primary_hex):
    """A deliberate, visually substantial section-divider page for each module — large,
    centered title, not a small heading sitting alone looking like a rendering mistake."""
    for _ in range(4):
        doc.add_paragraph()

    code_p = doc.add_paragraph()
    code_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    code_run = code_p.add_run(module.get("module_code", ""))
    code_run.font.size = Pt(16)
    code_run.font.color.rgb = primary
    code_run.italic = True

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(module.get("title", ""))
    title_run.bold = True
    title_run.font.size = Pt(36)
    title_run.font.color.rgb = primary

    rule_p = doc.add_paragraph()
    rule_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_bottom_border(rule_p, primary_hex, size="12")

    doc.add_page_break()


def _add_workbook_topic_section(doc, module, topic, primary, primary_hex, secondary, add_trailing_break=True):
    """Per-topic active-learning exercises — key terms, self-check ratings drawn from the
    topic's real assessment_criteria, notes space, and a takeaway line."""
    topic_code = topic.get("topic_code", "")
    title = topic.get("title", "")

    heading = doc.add_paragraph()
    h_run = heading.add_run(f"{topic_code}: {title}")
    h_run.bold = True
    h_run.font.size = Pt(15)
    h_run.font.color.rgb = primary
    _add_bottom_border(heading, primary_hex, size="8")

    elements = topic.get("elements") or []
    if elements:
        ref_p = doc.add_paragraph()
        ref_run = ref_p.add_run("Topic Elements Covered:")
        ref_run.bold = True
        ref_run.italic = True
        ref_run.font.size = Pt(10)
        ref_run.font.color.rgb = secondary
        for el in elements:
            code = _element_code(el)
            text = _element_text(el)
            label = f"{code}: {text}" if code else text
            item_p = doc.add_paragraph(label, style="List Bullet")
            for run in item_p.runs:
                run.font.size = Pt(10)
                run.italic = True
        doc.add_paragraph()

    doc.add_paragraph("Key Terms").runs[0].bold = True
    doc.add_paragraph("Write down any new or unfamiliar terms from this topic, in your own words.")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Term"
    table.rows[0].cells[1].text = "My Definition"
    for c in table.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True
    for _ in range(5):
        table.add_row()

    doc.add_paragraph()
    doc.add_paragraph("Self-Check").runs[0].bold = True
    criteria = topic.get("assessment_criteria") or []
    if criteria:
        doc.add_paragraph("Rate yourself honestly on each of the following:")
        sc_table = doc.add_table(rows=1, cols=4)
        sc_table.style = "Table Grid"
        hdr = sc_table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Can you...", "Not Yet", "Getting There", "Confident"
        for c in hdr:
            c.paragraphs[0].runs[0].bold = True
        for criterion in criteria:
            row = sc_table.add_row().cells
            row[0].text = criterion
    else:
        doc.add_paragraph(
            "No specific criteria were extracted for this topic — use the space below to note "
            "your own confidence level after the session."
        )
        doc.add_paragraph("_" * 100)

    doc.add_paragraph()
    doc.add_paragraph("My Notes").runs[0].bold = True
    for _ in range(6):
        doc.add_paragraph("_" * 100).paragraph_format.space_after = Pt(6)

    doc.add_paragraph()
    takeaway_p = doc.add_paragraph()
    takeaway_p.add_run("One Thing I Want to Remember From This Topic:").bold = True
    doc.add_paragraph("_" * 100)

    if add_trailing_break:
        doc.add_page_break()


def build_qcto_km_learner_workbook_docx(title: str, syllabus_content: dict, organization_name: str = None,
                                         logo_bytes: bytes = None, brand_colors: dict = None,
                                         accreditation_info: dict = None, job_id: str = None) -> BytesIO:
    accreditation_info = dict(accreditation_info or {})
    if syllabus_content.get("qualification_code"):
        accreditation_info.setdefault("qualification_code", syllabus_content.get("qualification_code"))
    brand_colors = brand_colors or {}
    primary_hex = brand_colors.get("primary", DEFAULT_PRIMARY).lstrip("#") if brand_colors.get("primary") else DEFAULT_PRIMARY
    secondary_hex = brand_colors.get("secondary", DEFAULT_SECONDARY).lstrip("#") if brand_colors.get("secondary") else DEFAULT_SECONDARY
    primary = _hex_to_rgb(primary_hex, DEFAULT_PRIMARY)
    secondary = _hex_to_rgb(secondary_hex, DEFAULT_SECONDARY)
    qualification_title = syllabus_content.get("qualification_title", "") or title

    km_modules = [m for m in syllabus_content.get("modules", []) if m.get("module_type") == "KM"]

    doc = Document()
    _build_branded_cover(doc, qualification_title, "KM Learner Workbook", organization_name, logo_bytes, primary, primary_hex, secondary)

    _add_workbook_intro(doc, primary, primary_hex)

    for m_index, module in enumerate(km_modules):
        _add_module_divider(doc, module, primary, primary_hex)

        topics = module.get("topics", [])
        for t_index, topic in enumerate(topics):
            is_very_last = (m_index == len(km_modules) - 1) and (t_index == len(topics) - 1)
            _add_workbook_topic_section(doc, module, topic, primary, primary_hex, secondary, add_trailing_break=not is_very_last)

    _add_branded_header_footer(doc, logo_bytes=logo_bytes, qualification_name=qualification_title, organization_name=organization_name, primary_hex=primary_hex, accreditation_info=accreditation_info, document_label="KM Learner Workbook")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_qcto_km_learner_workbook_docx_adapter(title, units, organization_name=None, seta=None,
                                                  nqf_level=None, logo_bytes=None, brand_colors=None,
                                                  accreditation_info=None, job_id=None, **kwargs):
    """Adapter matching the standard DOCUMENT_BUILDERS call signature."""
    syllabus_content = {"modules": units} if isinstance(units, list) else (units or {"modules": []})
    return build_qcto_km_learner_workbook_docx(
        title=title,
        syllabus_content=syllabus_content,
        organization_name=organization_name,
        logo_bytes=logo_bytes,
        brand_colors=brand_colors,
        accreditation_info=accreditation_info,
        job_id=job_id,
    )
